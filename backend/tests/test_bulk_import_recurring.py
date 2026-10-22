"""The Recurrentes block of the bulk import: subscriptions, services and recurring incomes."""

from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from tests.test_debts import _create_account, _register_and_login

pytestmark = pytest.mark.asyncio

REC_COL = 37  # AK; AP Categoría, AQ Tipo, AR Día del mes, AS Notas


def _build_upload(rows: list[dict], expenses: list[dict] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transacciones"
    ws["A2"] = "Monto"
    for offset, header in enumerate(["Nombre", "Monto", "Frecuencia", "Siguiente cobro", "Cuenta"]):
        ws.cell(row=2, column=REC_COL + offset, value=header)
    for col, header in ((42, "Categoría"), (43, "Tipo"), (44, "Día del mes"), (45, "Notas")):
        ws.cell(row=2, column=col, value=header)
    for row_num, row in enumerate(rows, start=3):
        for offset, key in enumerate(("name", "amount", "frequency", "next", "account")):
            ws.cell(row=row_num, column=REC_COL + offset, value=row.get(key))
        for col, key in ((42, "category"), (43, "type"), (44, "day"), (45, "notes")):
            if key in row:
                ws.cell(row=row_num, column=col, value=row[key])
    # expenses: [{amount, description, date, account, category, recurring}] in G-K; link in AT
    ws.cell(row=2, column=46, value="Recurrente (gastos)")
    for row_num, row in enumerate(expenses or [], start=3):
        for offset, key in enumerate(("amount", "description", "date", "account", "category")):
            ws.cell(row=row_num, column=7 + offset, value=row.get(key))
        if row.get("recurring") is not None:
            ws.cell(row=row_num, column=46, value=row["recurring"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _post(client, headers, path, content):
    return await client.post(
        f"/api/v1/bulk-import/{path}",
        headers=headers,
        files={"file": ("recurrentes.xlsx", content, "application/octet-stream")},
    )


async def _category(client, headers, type_):
    response = await client.get(f"/api/v1/categories?type={type_}", headers=headers)
    return response.json()["data"][0]["name"]


async def _recurring(client, headers):
    response = await client.get("/api/v1/recurring-items", headers=headers)
    return response.json()["data"]


async def test_template_has_the_recurring_block(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    ws = load_workbook(BytesIO(response.content))["Transacciones"]
    assert ws["AK1"].value == "Recurrentes"
    assert [ws.cell(row=2, column=c).value for c in range(37, 46)] == [
        "Nombre",
        "Monto",
        "Frecuencia",
        "Siguiente cobro",
        "Cuenta",
        "Categoría",
        "Tipo",
        "Día del mes",
        "Notas",
    ]
    assert ws["AT2"].value == "Recurrente (gastos)"
    assert ws["AU2"].value == "Recurrente (ingresos)"


async def test_upload_creates_the_recurring_items(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    santander = await _create_account(client, headers, name="Santander", initial_balance="1000")
    expense = await _category(client, headers, "expense")
    income = await _category(client, headers, "income")

    content = _build_upload(
        [
            {
                "name": "Seguro Santander - Blindaje total",
                "amount": 186.38,
                "frequency": "Mensual",
                "next": "2026-11-12",
                "account": "Santander",
                "category": expense,
                "type": "Servicio",
                "day": 12,
                "notes": "Póliza anti robo",
            },
            {
                "name": "Spotify",
                "amount": 239,
                "frequency": "mensual",
                "next": "2026-11-06",
                "account": "Santander",
                "category": expense,
            },
            {
                "name": "Honorarios fijos",
                "amount": 2000,
                "frequency": "Semanal",
                "next": "2026-10-25",
                "account": "Santander",
                "category": income,
                "type": "Ingreso",
            },
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 3, "errors": []}

    by_name = {i["name"]: i for i in await _recurring(client, headers)}
    seguro = by_name["Seguro Santander - Blindaje total"]
    assert (seguro["item_type"], seguro["frequency"], seguro["frequency_day"]) == (
        "service",
        "monthly",
        12,
    )
    assert seguro["amount"] == "186.38"
    assert seguro["next_date"] == "2026-11-12"
    assert seguro["account_id"] == santander
    assert seguro["notes"] == "Póliza anti robo"
    assert seguro["alert_urgency"] == "high"  # what a service gets by default
    # the type is optional: a subscription by default
    assert by_name["Spotify"]["item_type"] == "subscription"
    assert by_name["Honorarios fijos"]["item_type"] == "income"
    assert by_name["Honorarios fijos"]["frequency"] == "weekly"


async def test_preview_counts_the_recurring_items_and_keeps_them_out_of_the_date_range(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")
    expense = await _category(client, headers, "expense")

    content = _build_upload(
        [
            {
                "name": "Spotify",
                "amount": 239,
                "frequency": "Mensual",
                "next": "2026-11-06",
                "account": "Santander",
                "category": expense,
            },
            {
                "name": "Netflix",
                "amount": 200,
                "frequency": "Mensual",
                "next": "2026-11-09",
                "account": "Santander",
                "category": expense,
            },
        ]
    )
    response = await _post(client, headers, "preview", content)
    data = response.json()["data"]
    assert data["recurring"]["count"] == 2
    assert float(data["recurring"]["total"]) == 439
    assert data["valid_rows"] == 2
    assert data["date_from"] is None and data["date_to"] is None
    assert await _recurring(client, headers) == []  # a preview writes nothing


async def test_rows_that_do_not_fit_are_reported_and_the_rest_loads(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")
    expense = await _category(client, headers, "expense")
    income = await _category(client, headers, "income")

    ok = {
        "amount": 100,
        "frequency": "Mensual",
        "next": "2026-11-01",
        "account": "Santander",
        "category": expense,
    }
    content = _build_upload(
        [
            {**ok, "name": "Buena"},
            {**ok, "name": "Frecuencia mala", "frequency": "Cada rato"},
            {**ok, "name": "Sin fecha", "next": None},
            {**ok, "name": "Cuenta rara", "account": "No existe"},
            {**ok, "name": "Categoría rara", "category": "No existe"},
            {**ok, "name": "Ingreso con categoría de gasto", "type": "Ingreso"},
            {**ok, "name": "Gasto con categoría de ingreso", "category": income},
            {**ok, "name": "Tipo raro", "type": "Hipoteca"},
            {**ok, "name": "Día raro", "day": 40},
            {**ok, "name": "Monto cero", "amount": 0},
            {**ok, "name": "Buena"},  # repeated in the same file
        ]
    )
    response = await _post(client, headers, "upload", content)
    data = response.json()["data"]
    assert data["created"] == 1
    assert sorted(e["row"] for e in data["errors"]) == list(range(4, 14))
    assert [i["name"] for i in await _recurring(client, headers)] == ["Buena"]


async def test_an_active_recurring_item_with_the_same_name_and_account_is_not_duplicated(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")
    expense = await _category(client, headers, "expense")
    row = {
        "name": "Seguro",
        "amount": 175.05,
        "frequency": "Mensual",
        "next": "2026-11-12",
        "account": "Santander",
        "category": expense,
    }

    first = await _post(client, headers, "upload", _build_upload([row]))
    assert first.json()["data"]["created"] == 1
    again = await _post(client, headers, "upload", _build_upload([row]))
    data = again.json()["data"]
    assert data["created"] == 0
    assert "Ya tienes un recurrente activo" in data["errors"][0]["reason"]
    assert len(await _recurring(client, headers)) == 1


async def _expense(client, headers, description):
    response = await client.get("/api/v1/transactions?per_page=100", headers=headers)
    return [t for t in response.json()["data"] if t["description"] == description]


async def test_expenses_can_point_to_a_recurring_item_of_the_same_file(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="1000")
    expense = await _category(client, headers, "expense")

    content = _build_upload(
        [
            {
                "name": "Seguro",
                "amount": 175.05,
                "frequency": "Mensual",
                "next": "2026-11-12",
                "account": "Santander",
                "category": expense,
            }
        ],
        expenses=[
            # the expense comes BEFORE its recurring item in the file; the link still works
            {
                "amount": 175.05,
                "description": "Pago seguro 1",
                "date": "2026-09-12",
                "account": "Santander",
                "category": expense,
                "recurring": "Seguro",
            },
            {
                "amount": 175.05,
                "description": "Pago seguro 2",
                "date": "2026-10-12",
                "account": "Santander",
                "category": expense,
                "recurring": "seguro",
            },
            {
                "amount": 50,
                "description": "Gasto normal",
                "date": "2026-10-13",
                "account": "Santander",
                "category": expense,
            },
        ],
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 4, "errors": []}

    linked = await client.get(
        "/api/v1/transactions?is_recurring=true&per_page=100", headers=headers
    )
    assert sorted(t["description"] for t in linked.json()["data"]) == [
        "Pago seguro 1",
        "Pago seguro 2",
    ]
    plain = await _expense(client, headers, "Gasto normal")
    assert plain[0]["is_recurring"] is False
    # the recurring item itself is untouched: its next charge stays where it was
    (item,) = await _recurring(client, headers)
    assert item["next_date"] == "2026-11-12"


async def test_expenses_can_point_to_a_recurring_item_that_already_exists(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="1000")
    expense = await _category(client, headers, "expense")
    row = {
        "name": "Seguro",
        "amount": 175.05,
        "frequency": "Mensual",
        "next": "2026-11-12",
        "account": "Santander",
        "category": expense,
    }
    await _post(client, headers, "upload", _build_upload([row]))

    content = _build_upload(
        [],
        expenses=[
            {
                "amount": 175.05,
                "description": "Pago seguro viejo",
                "date": "2026-08-12",
                "account": "Santander",
                "category": expense,
                "recurring": "Seguro",
            }
        ],
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 1, "errors": []}
    linked = await client.get(
        "/api/v1/transactions?is_recurring=true&per_page=100", headers=headers
    )
    assert [t["description"] for t in linked.json()["data"]] == ["Pago seguro viejo"]


async def test_a_link_to_an_unknown_or_mismatched_recurring_item_is_reported(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="1000")
    expense = await _category(client, headers, "expense")
    income = await _category(client, headers, "income")
    sueldo = {
        "name": "Sueldo",
        "amount": 100,
        "frequency": "Mensual",
        "next": "2026-11-01",
        "account": "Santander",
        "category": income,
        "type": "Ingreso",
    }
    await _post(client, headers, "upload", _build_upload([sueldo]))

    content = _build_upload(
        [],
        expenses=[
            {
                "amount": 10,
                "description": "A",
                "date": "2026-10-01",
                "account": "Santander",
                "category": expense,
                "recurring": "No existe",
            },
            {
                "amount": 10,
                "description": "B",
                "date": "2026-10-01",
                "account": "Santander",
                "category": expense,
                "recurring": "Sueldo",
            },
            {
                "amount": 10,
                "description": "C",
                "date": "2026-10-01",
                "account": "Santander",
                "category": expense,
            },
        ],
    )
    response = await _post(client, headers, "upload", content)
    data = response.json()["data"]
    assert data["created"] == 1
    assert sorted(e["row"] for e in data["errors"]) == [3, 4]
    assert "No existe el recurrente" in data["errors"][0]["reason"]
    assert "es de ingresos" in data["errors"][1]["reason"]
