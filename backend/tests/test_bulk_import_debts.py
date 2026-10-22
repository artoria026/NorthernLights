"""The Deudas block of the bulk import: loans, repayments and opening balances by person."""

from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from tests.test_debts import _create_account, _register_and_login

pytestmark = pytest.mark.asyncio

DEBT_COL = 23  # W: Monto, Persona, Movimiento, Fecha, Cuenta, Concepto; AC (29): Dirección


def _build_debt_upload(rows: list[dict]) -> bytes:
    """rows: [{'amount', 'person', 'kind', 'date', 'account', 'concept', 'direction'}] in the
    debts block (W-AC), data starting from row 3."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transacciones"
    ws["A2"] = "Monto"
    for offset, header in enumerate(
        ["Monto", "Persona", "Movimiento", "Fecha", "Cuenta", "Concepto"]
    ):
        ws.cell(row=2, column=DEBT_COL + offset, value=header)
    ws.cell(row=2, column=29, value="Dirección")
    for row_num, row in enumerate(rows, start=3):
        ws.cell(row=row_num, column=DEBT_COL, value=row.get("amount"))
        ws.cell(row=row_num, column=DEBT_COL + 1, value=row.get("person"))
        ws.cell(row=row_num, column=DEBT_COL + 2, value=row.get("kind"))
        ws.cell(row=row_num, column=DEBT_COL + 3, value=row.get("date"))
        ws.cell(row=row_num, column=DEBT_COL + 4, value=row.get("account"))
        ws.cell(row=row_num, column=DEBT_COL + 5, value=row.get("concept"))
        if "direction" in row:
            ws.cell(row=row_num, column=29, value=row["direction"])
        # optional payment plan (AD-AH): amount, frequency, day, next due date, account
        for offset, key in enumerate(("plan", "frequency", "day", "next", "plan_account", "type")):
            if key in row:
                ws.cell(row=row_num, column=30 + offset, value=row[key])
        # terms of a debt with interest (AV-AX): lent, total to pay, yearly rate
        for offset, key in enumerate(("lent", "total", "rate")):
            if key in row:
                ws.cell(row=row_num, column=48 + offset, value=row[key])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _post(client, headers, path, content):
    return await client.post(
        f"/api/v1/bulk-import/{path}",
        headers=headers,
        files={"file": ("deudas.xlsx", content, "application/octet-stream")},
    )


async def _debts(client, headers):
    return (await client.get("/api/v1/debts", headers=headers)).json()["data"]


async def _balance(client, headers, account_id):
    response = await client.get(f"/api/v1/accounts/{account_id}", headers=headers)
    return response.json()["data"]["balance"]


async def test_template_has_the_debts_block(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    ws = load_workbook(BytesIO(response.content))["Transacciones"]
    assert ws["W1"].value == "Deudas"
    assert [ws.cell(row=2, column=c).value for c in range(23, 36)] == [
        "Monto",
        "Persona",
        "Movimiento",
        "Fecha",
        "Cuenta",
        "Concepto",
        "Dirección",
        "Pago periódico",
        "Frecuencia",
        "Día de pago",
        "Siguiente pago",
        "Cuenta de pago",
        "Tipo",
    ]


async def test_loans_and_repayments_build_a_persons_debt_in_date_order(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    santander = await _create_account(client, headers, name="Santander", initial_balance="5000")
    efectivo = await _create_account(client, headers, name="Efectivo", initial_balance="0")

    # The repayment comes FIRST in the file, but it is dated after both loans.
    content = _build_debt_upload(
        [
            {
                "amount": 300,
                "person": "Emanuel",
                "kind": "Abono",
                "date": "2026-10-10",
                "account": "Efectivo",
                "concept": "Abono en efectivo",
            },
            {
                "amount": 1000,
                "person": "emanuel",
                "kind": "Préstamo",
                "date": "2026-10-05",
                "account": "Santander",
                "concept": "Reparación del coche",
            },
            {
                "amount": 500,
                "person": "Emanuel",
                "kind": "Prestamo",
                "date": "2026-10-01",
                "account": "Santander",
                "concept": "Cena de cumpleaños",
            },
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 3, "errors": []}

    debts = await _debts(client, headers)
    assert len(debts) == 1
    assert debts[0]["name"] == "Emanuel"
    assert debts[0]["direction"] == "owed_to_me"
    assert debts[0]["current_balance"] == "1200.00"
    assert debts[0]["total_amount"] == "1500.00"

    movements = (
        await client.get(f"/api/v1/debts/{debts[0]['id']}/movements", headers=headers)
    ).json()["data"]
    assert [
        (m["kind"], m["delta"], m["balance_after"], m["concept"], m["account_name"])
        for m in movements
    ] == [
        ("loan", "500.00", "500.00", "Cena de cumpleaños", "Santander"),
        ("loan", "1000.00", "1500.00", "Reparación del coche", "Santander"),
        ("payment", "-300.00", "1200.00", "Abono en efectivo", "Efectivo"),
    ]
    # the real accounts moved: 1,500 lent out of Santander, 300 collected in cash
    assert await _balance(client, headers, santander) == "3500.00"
    assert await _balance(client, headers, efectivo) == "300.00"


async def test_opening_balance_needs_no_account_and_direction_debo_is_a_debt_of_mine(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    santander = await _create_account(client, headers, name="Santander", initial_balance="1000")

    content = _build_debt_upload(
        [
            {
                "amount": 15000,
                "person": "Franco Andrade",
                "kind": "Saldo inicial",
                "date": "2026-01-01",
                "concept": "Lo que me debía antes de la app",
            },
            {
                "amount": 800,
                "person": "Mireya",
                "kind": "Préstamo",
                "date": "2026-10-02",
                "account": "Santander",
                "concept": "Para la renta",
                "direction": "Debo",
            },
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 2, "errors": []}

    by_name = {d["name"]: d for d in await _debts(client, headers)}
    assert by_name["Franco Andrade"]["direction"] == "owed_to_me"
    assert by_name["Franco Andrade"]["current_balance"] == "15000.00"
    assert by_name["Mireya"]["direction"] == "owed_by_me"
    # the opening balance moved no money; what I borrowed from Mireya landed in Santander
    assert await _balance(client, headers, santander) == "1800.00"
    movements = (
        await client.get(
            f"/api/v1/debts/{by_name['Franco Andrade']['id']}/movements", headers=headers
        )
    ).json()["data"]
    assert [(m["kind"], m["account_id"]) for m in movements] == [("opening", None)]


async def test_rows_that_do_not_fit_are_reported_and_the_rest_loads(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="5000")

    content = _build_debt_upload(
        [
            {
                "amount": 500,
                "person": "Emanuel",
                "kind": "Préstamo",
                "date": "2026-10-01",
                "account": "Santander",
                "concept": "Cena",
            },
            {
                "amount": 50,
                "person": "Nadie",
                "kind": "Abono",
                "date": "2026-10-02",
                "account": "Santander",
            },  # nobody owes me yet
            {
                "amount": 900,
                "person": "Emanuel",
                "kind": "Abono",
                "date": "2026-10-03",
                "account": "Santander",
            },  # more than he owes
            {
                "amount": 100,
                "person": "Emanuel",
                "kind": "Préstamo",
                "date": "2026-10-04",
            },  # no account
            {
                "amount": 100,
                "person": "Emanuel",
                "kind": "Préstamo",
                "date": "2026-10-04",
                "account": "Inexistente",
            },
            {
                "amount": 100,
                "person": "Emanuel",
                "kind": "Regalo",
                "date": "2026-10-04",
                "account": "Santander",
            },  # not a movement
            {"amount": 700, "person": "Emanuel", "kind": "Saldo inicial", "date": "2026-10-05"},
        ]
    )
    response = await _post(client, headers, "upload", content)
    data = response.json()["data"]
    assert data["created"] == 1
    reasons = {e["row"]: e["reason"] for e in data["errors"]}
    assert sorted(reasons) == [4, 5, 6, 7, 8, 9]
    assert "No hay una deuda abierta con Nadie" in reasons[4]
    assert "mayor que lo que Emanuel debe" in reasons[5]
    assert "Falta la cuenta" in reasons[6]
    assert "'Inexistente' no existe" in reasons[7]
    assert "Movimiento inválido" in reasons[8]
    assert "Ya existe una deuda abierta con Emanuel" in reasons[9]
    assert all(e["block"] == "Deuda" for e in data["errors"])
    assert (await _debts(client, headers))[0]["current_balance"] == "500.00"


async def test_a_loan_adds_to_a_debt_that_already_exists_and_reopens_a_paid_one(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    santander = await _create_account(client, headers, name="Santander", initial_balance="5000")
    created = await client.post(
        "/api/v1/debts",
        headers=headers,
        json={
            "name": "Emanuel",
            "type": "informal",
            "direction": "owed_to_me",
            "total_amount": "500.00",
            "funding_account_id": santander,
            "start_date": "2026-09-01",
        },
    )
    debt_id = created.json()["data"]["id"]
    await client.post(
        f"/api/v1/debts/{debt_id}/payments",
        headers=headers,
        json={"account_id": santander, "amount": "500.00", "date": "2026-09-15"},
    )
    assert (await client.get(f"/api/v1/debts/{debt_id}", headers=headers)).json()["data"][
        "status"
    ] == "completed"

    content = _build_debt_upload(
        [
            {
                "amount": 250,
                "person": "Emanuel",
                "kind": "Préstamo",
                "date": "2026-10-01",
                "account": "Santander",
                "concept": "Otra vez",
            }
        ]
    )
    assert (await _post(client, headers, "upload", content)).json()["data"] == {
        "created": 1,
        "errors": [],
    }

    debts = await _debts(client, headers)
    assert len(debts) == 1  # the same debt, not a second one
    assert debts[0]["status"] == "active"
    assert debts[0]["current_balance"] == "250.00"


async def test_preview_summarizes_the_debts_without_saving_and_warns_about_repeats(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    santander = await _create_account(client, headers, name="Santander", initial_balance="5000")
    content = _build_debt_upload(
        [
            {
                "amount": 500,
                "person": "Emanuel",
                "kind": "Préstamo",
                "date": "2026-10-01",
                "account": "Santander",
                "concept": "Cena",
            },
            {
                "amount": 1000,
                "person": "Emanuel",
                "kind": "Préstamo",
                "date": "2026-10-05",
                "account": "Santander",
                "concept": "Coche",
            },
            {
                "amount": 300,
                "person": "Emanuel",
                "kind": "Abono",
                "date": "2026-10-10",
                "account": "Santander",
            },
            {"amount": 15000, "person": "Franco", "kind": "Saldo inicial", "date": "2026-01-01"},
        ]
    )
    data = (await _post(client, headers, "preview", content)).json()["data"]
    assert data["valid_rows"] == 4
    assert data["debts"] == {
        "loans": {"count": 2, "total": "1500"},
        "payments": {"count": 1, "total": "300"},
        "openings": {"count": 1, "total": "15000"},
        "people": 2,
    }
    assert data["income"]["count"] == 0 and data["expense"]["count"] == 0
    assert data["possible_duplicates"]["count"] == 0
    # nothing was saved by the preview
    assert await _debts(client, headers) == []
    assert await _balance(client, headers, santander) == "5000.00"

    await _post(client, headers, "upload", content)
    again = (await _post(client, headers, "preview", content)).json()["data"]
    # a repeat of the file: the loans and the repayment look already loaded (the opening
    # balance fails because Franco's debt is open now)
    assert again["possible_duplicates"]["count"] >= 3


async def test_payment_plan_is_set_after_the_repayments_so_the_next_date_is_kept(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    santander = await _create_account(client, headers, name="Santander", initial_balance="9000")

    # The plan sits on the opening row, but two repayments come after it in the file.
    content = _build_debt_upload(
        [
            {
                "amount": 5000,
                "person": "Pavimento calle Cubelita",
                "kind": "Saldo inicial",
                "date": "2026-01-01",
                "direction": "Debo",
                "plan": 1000,
                "frequency": "Mensual",
                "day": 17,
                "next": "2026-11-17",
                "plan_account": "Santander",
            },
            {
                "amount": 1000,
                "person": "Pavimento calle Cubelita",
                "kind": "Abono",
                "date": "2026-02-17",
                "account": "Santander",
                "direction": "Debo",
            },
            {
                "amount": 1000,
                "person": "Pavimento calle Cubelita",
                "kind": "Abono",
                "date": "2026-03-17",
                "account": "Santander",
                "direction": "Debo",
            },
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 3, "errors": []}

    (debt,) = await _debts(client, headers)
    assert debt["current_balance"] == "3000.00"
    assert debt["payment_amount"] == "1000.00"
    assert debt["payment_frequency"] == "monthly"
    assert debt["payment_day"] == 17
    assert debt["next_payment_date"] == "2026-11-17"
    assert debt["payment_source_account_id"] == santander
    movements = (await client.get(f"/api/v1/debts/{debt['id']}/movements", headers=headers)).json()[
        "data"
    ]
    assert [m["kind"] for m in movements] == ["opening", "payment", "payment"]


async def test_payment_plan_defaults_to_monthly_and_is_optional_per_field(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="0")

    content = _build_debt_upload(
        [
            {
                "amount": 900,
                "person": "Educación Física",
                "kind": "Saldo inicial",
                "date": "2026-10-21",
                "direction": "Debo",
                "plan": 300,
            }
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 1, "errors": []}

    (debt,) = await _debts(client, headers)
    assert debt["payment_amount"] == "300.00"
    assert debt["payment_frequency"] == "monthly"
    assert debt["next_payment_date"] is None
    assert debt["payment_source_account_id"] is None


async def test_invalid_payment_plans_are_reported_without_loading_the_row(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="0")
    await _create_account(client, headers, name="Tarjeta", type="liability", subtype="credit_card")

    base = {"amount": 900, "kind": "Saldo inicial", "date": "2026-10-21", "direction": "Debo"}
    content = _build_debt_upload(
        [
            {**base, "person": "A", "frequency": "Mensual"},  # a plan without its payment
            {**base, "person": "B", "plan": 100, "frequency": "Anual"},
            {**base, "person": "C", "plan": 100, "day": 40},
            {**base, "person": "D", "plan": 100, "next": "no es fecha"},
            {**base, "person": "E", "plan": 100, "plan_account": "No existe"},
            {**base, "person": "F", "plan": 100, "plan_account": "Tarjeta"},
        ]
    )
    response = await _post(client, headers, "upload", content)
    data = response.json()["data"]
    assert data["created"] == 0
    assert sorted(e["row"] for e in data["errors"]) == [3, 4, 5, 6, 7, 8]
    assert await _debts(client, headers) == []


async def test_debt_type_column_creates_the_debt_with_that_kind_and_later_rows_find_it(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="5000")

    # The kind sits on the opening row only; the repayment finds the debt by its name.
    content = _build_debt_upload(
        [
            {
                "amount": 28200,
                "person": "Pavimento calle Cubelita",
                "kind": "Saldo inicial",
                "date": "2025-07-11",
                "direction": "Debo",
                "type": "Crédito cívico",
            },
            {
                "amount": 7500,
                "person": "Pavimento calle Cubelita",
                "kind": "Abono",
                "date": "2025-07-11",
                "account": "Santander",
                "direction": "Debo",
            },
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 2, "errors": []}

    (debt,) = await _debts(client, headers)
    assert debt["type"] == "civic"
    assert debt["current_balance"] == "20700.00"


async def test_debt_type_is_informal_by_default_and_an_unknown_one_is_reported(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="0")

    base = {"amount": 100, "kind": "Saldo inicial", "date": "2026-10-21"}
    content = _build_debt_upload(
        [
            {**base, "person": "Emanuel"},
            {**base, "person": "Banco", "direction": "Debo", "type": "Hipoteca"},
        ]
    )
    response = await _post(client, headers, "upload", content)
    data = response.json()["data"]
    assert data["created"] == 1
    assert [e["row"] for e in data["errors"]] == [4]
    assert [d["type"] for d in await _debts(client, headers)] == ["informal"]


async def test_template_has_the_debt_terms_columns(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    ws = load_workbook(BytesIO(response.content))["Transacciones"]
    assert [ws.cell(row=2, column=c).value for c in (48, 49, 50)] == [
        "Monto prestado",
        "Total a pagar",
        "Tasa anual (%)",
    ]


async def test_a_debt_with_interest_keeps_what_was_lent_and_the_total_to_pay(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="5000")

    # 50,000 lent, 98,700 to pay in all; when the record starts 76,000 is still owed. The weekly
    # payments include interest, so together they are more than the principal.
    content = _build_debt_upload(
        [
            {
                "amount": 76000,
                "person": "Préstamo nómina",
                "kind": "Saldo inicial",
                "date": "2025-05-15",
                "direction": "Debo",
                "type": "Crédito de nómina",
                "lent": 50000,
                "total": 98700,
                "rate": "47.5%",
            },
            {
                "amount": 600,
                "person": "Préstamo nómina",
                "kind": "Abono",
                "date": "2025-05-15",
                "account": "Santander",
                "direction": "Debo",
            },
            {
                "amount": 75400,
                "person": "Préstamo nómina",
                "kind": "Abono",
                "date": "2026-09-04",
                "account": "Santander",
                "direction": "Debo",
            },
        ]
    )
    response = await _post(client, headers, "upload", content)
    assert response.json()["data"] == {"created": 3, "errors": []}

    listing = (await client.get("/api/v1/debts?direction=owed_by_me", headers=headers)).json()[
        "data"
    ]
    (debt,) = listing
    assert debt["type"] == "payroll_loan"
    assert debt["original_amount"] == "50000.00"
    assert debt["total_amount"] == "98700.00"
    assert debt["current_balance"] == "0.00"
    assert debt["status"] == "completed"
    assert debt["interest_rate"] == "0.4750"


async def test_debt_terms_are_checked(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander", initial_balance="5000")

    base = {"person": "Crédito", "date": "2025-05-15", "direction": "Debo"}
    content = _build_debt_upload(
        [
            {
                **base,
                "amount": 1000,
                "kind": "Saldo inicial",
                "total": 900,
            },  # total below the balance
            {**base, "person": "B", "amount": 1000, "kind": "Saldo inicial", "rate": 150},
            {**base, "person": "C", "amount": 1000, "kind": "Saldo inicial", "lent": 0},
            {
                **base,
                "person": "D",
                "amount": 100,
                "kind": "Préstamo",
                "account": "Santander",
                "total": 200,
            },
            {**base, "person": "E", "amount": 1000, "kind": "Saldo inicial", "total": 1500},
        ]
    )
    response = await _post(client, headers, "upload", content)
    data = response.json()["data"]
    assert data["created"] == 1
    assert sorted(e["row"] for e in data["errors"]) == [3, 4, 5, 6]
