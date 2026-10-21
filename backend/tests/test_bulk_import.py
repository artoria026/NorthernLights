import uuid
from decimal import Decimal
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

pytestmark = pytest.mark.asyncio


async def _register_and_login(client: AsyncClient) -> str:
    email = f"{uuid.uuid4()}@example.com"
    await client.post(
        "/api/v1/auth/register",
        json={
            "email": email,
            "name": "Test",
            "password": "supersecret123",
            "accept_disclaimer": True,
        },
    )
    login = await client.post(
        "/api/v1/auth/login", json={"email": email, "password": "supersecret123"}
    )
    return login.json()["data"]["access_token"]


async def _create_account(client: AsyncClient, headers: dict, **overrides) -> str:
    payload = {"name": "Santander", "type": "asset", "subtype": "checking", "initial_balance": "0"}
    payload.update(overrides)
    response = await client.post("/api/v1/accounts", headers=headers, json=payload)
    assert response.status_code == 201
    return response.json()["data"]["id"]


def _build_upload(rows: list[dict]) -> bytes:
    """rows: [{'block': 'income'|'expense', 'amount', 'description', 'date',
    'account', 'category'}] -- builds an .xlsx with the same layout that
    build_template_workbook generates (Income in A-E, Expenses in G-K, data starting
    from row 3)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transacciones"
    ws["A2"] = "Monto"
    ws["G2"] = "Monto"

    row_num = 3
    for row in rows:
        start_col = {"income": "A", "expense": "G", "card_payment": "P"}[row["block"]]
        cols = [chr(ord(start_col) + i) for i in range(5)]
        ws[f"{cols[0]}{row_num}"] = row["amount"]
        ws[f"{cols[1]}{row_num}"] = row["description"]
        ws[f"{cols[2]}{row_num}"] = row["date"]
        ws[f"{cols[3]}{row_num}"] = row["account"]
        # in a card payment block the last column is the card, not a category
        ws[f"{cols[4]}{row_num}"] = row["category"]
        if "tags" in row:
            ws[f"L{row_num}"] = row["tags"]  # optional "Etiquetas" column of the Expenses block
        if "months" in row:
            ws[f"U{row_num}"] = row["months"]  # optional "Meses sin intereses" column
        if "notes" in row:
            # optional notes: "Notas" (M) for an expense, "Notas (ingresos)" (N) for an income
            ws[f"{'M' if row['block'] == 'expense' else 'N'}{row_num}"] = row["notes"]
        row_num += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_download_template_is_valid_xlsx_with_dropdowns(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    wb = load_workbook(BytesIO(response.content))
    assert "Transacciones" in wb.sheetnames
    assert "Listas" in wb.sheetnames
    ws = wb["Transacciones"]
    assert ws["A2"].value == "Monto"
    assert ws["G2"].value == "Monto"
    assert len(ws.data_validations.dataValidation) == 9


async def test_upload_creates_valid_rows_and_reports_row_errors(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers, name="Santander")

    content = _build_upload(
        [
            {
                "block": "income",
                "amount": 2807,
                "description": "Pago nomina",
                "date": "2026-08-01",
                "account": "Santander",
                "category": "Empleo principal",
            },
            {
                "block": "expense",
                "amount": 35,
                "description": "Burritos",
                "date": "2026-08-01",
                "account": "Santander",
                "category": "Comida y Bebidas",
            },
            {
                "block": "expense",
                "amount": 100,
                "description": "Cuenta invalida",
                "date": "2026-08-02",
                "account": "Cuenta que no existe",
                "category": "Comida y Bebidas",
            },
            {
                "block": "expense",
                "amount": 50,
                "description": "Categoria invalida",
                "date": "2026-08-02",
                "account": "Santander",
                "category": "Categoria que no existe",
            },
            {
                "block": "expense",
                "amount": -10,
                "description": "Monto invalido",
                "date": "2026-08-02",
                "account": "Santander",
                "category": "Comida y Bebidas",
            },
        ]
    )

    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={
            "file": (
                "carga.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["created"] == 2
    reasons = {e["reason"] for e in data["errors"]}
    assert any("no existe" in r and "Cuenta" in r for r in reasons)
    assert any("no existe" in r and "categor" in r.lower() for r in reasons)
    assert any("Monto inválido" in r for r in reasons)

    transactions = await client.get("/api/v1/transactions?per_page=10", headers=headers)
    descriptions = {t["description"] for t in transactions.json()["data"]}
    assert descriptions == {"Pago nomina", "Burritos"}


async def test_upload_missing_sheet_reports_clear_error(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}

    wb = Workbook()
    wb.active.title = "OtraHoja"
    buf = BytesIO()
    wb.save(buf)

    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={
            "file": (
                "malo.xlsx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["created"] == 0
    assert len(data["errors"]) == 1
    assert "Transacciones" in data["errors"][0]["reason"]


async def test_upload_rls_isolation(client: AsyncClient):
    token_a = await _register_and_login(client)
    token_b = await _register_and_login(client)
    headers_a = {"Authorization": f"Bearer {token_a}"}
    headers_b = {"Authorization": f"Bearer {token_b}"}
    await _create_account(client, headers_b, name="Cuenta de B")

    content = _build_upload(
        [
            {
                "block": "expense",
                "amount": 10,
                "description": "Gasto de A intentando usar cuenta de B",
                "date": "2026-08-01",
                "account": "Cuenta de B",
                "category": "Comida y Bebidas",
            }
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers_a,
        files={
            "file": (
                "carga.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    assert response.json()["data"]["created"] == 0  # A can't see B's account


async def _get_category_names(client: AsyncClient, headers: dict, type_: str) -> str:
    response = await client.get(f"/api/v1/categories?type={type_}", headers=headers)
    return response.json()["data"][0]["name"]


async def _create_tag(client: AsyncClient, headers: dict, name: str) -> str:
    response = await client.post("/api/v1/tags", headers=headers, json={"name": name})
    assert response.status_code == 201
    return response.json()["data"]["id"]


async def test_template_has_the_tags_column_and_suggests_the_users_tags(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    await _create_tag(client, headers, "Reparación choque")

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    wb = load_workbook(BytesIO(response.content))
    ws = wb["Transacciones"]
    assert ws["L2"].value == "Etiquetas"
    assert len(ws.data_validations.dataValidation) == 10
    assert wb["Listas"]["D1"].value == "Reparación choque"


async def test_upload_puts_the_tags_of_each_expense(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    category = await _get_category_names(client, headers, "expense")
    await _create_tag(client, headers, "Viaje Mty")
    await _create_tag(client, headers, "Con pareja")

    content = _build_upload(
        [
            {
                "block": "expense",
                "amount": 100,
                "description": "Cena",
                "date": "2026-10-12",
                "account": "Santander",
                "category": category,
                "tags": "viaje mty, Con pareja",
            },
            {
                "block": "expense",
                "amount": 50,
                "description": "Taxi",
                "date": "2026-10-12",
                "account": "Santander",
                "category": category,
                "tags": " Viaje Mty ; viaje mty ",
            },
            {
                "block": "expense",
                "amount": 30,
                "description": "Sin etiquetas",
                "date": "2026-10-12",
                "account": "Santander",
                "category": category,
            },
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    assert response.json()["data"] == {"created": 3, "errors": []}

    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    tags = {t["description"]: sorted(x["name"] for x in t["tags"]) for t in listed.json()["data"]}
    assert tags == {"Cena": ["Con pareja", "Viaje Mty"], "Taxi": ["Viaje Mty"], "Sin etiquetas": []}


async def test_upload_rejects_the_row_with_an_unknown_tag_and_loads_the_rest(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    category = await _get_category_names(client, headers, "expense")
    await _create_tag(client, headers, "Viaje Mty")

    content = _build_upload(
        [
            {
                "block": "expense",
                "amount": 100,
                "description": "Bien",
                "date": "2026-10-12",
                "account": "Santander",
                "category": category,
                "tags": "Viaje Mty",
            },
            {
                "block": "expense",
                "amount": 40,
                "description": "Typo",
                "date": "2026-10-12",
                "account": "Santander",
                "category": category,
                "tags": "Viaje Mtyy",
            },
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    data = response.json()["data"]
    assert data["created"] == 1
    assert len(data["errors"]) == 1
    assert data["errors"][0]["row"] == 4 and "Viaje Mtyy" in data["errors"][0]["reason"]
    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    assert [t["description"] for t in listed.json()["data"]] == ["Bien"]


async def test_tags_of_one_user_are_not_usable_by_another(client: AsyncClient):
    other = {"Authorization": f"Bearer {await _register_and_login(client)}"}
    await _create_tag(client, other, "Secreta")
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    category = await _get_category_names(client, headers, "expense")

    content = _build_upload(
        [
            {
                "block": "expense",
                "amount": 10,
                "description": "Ajena",
                "date": "2026-10-12",
                "account": "Santander",
                "category": category,
                "tags": "Secreta",
            }
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    assert response.json()["data"]["created"] == 0
    assert "Secreta" in response.json()["data"]["errors"][0]["reason"]


async def _preview(client: AsyncClient, headers: dict, content: bytes) -> dict:
    response = await client.post(
        "/api/v1/bulk-import/preview",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


async def test_preview_summarizes_the_file_without_saving_anything(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    income = await _get_category_names(client, headers, "income")
    expense = await _get_category_names(client, headers, "expense")
    await _create_tag(client, headers, "Viaje Mty")

    content = _build_upload(
        [
            {
                "block": "income",
                "amount": 1000,
                "description": "Nomina",
                "date": "2026-10-01",
                "account": "Santander",
                "category": income,
            },
            {
                "block": "expense",
                "amount": 100,
                "description": "Cena",
                "date": "2026-10-12",
                "account": "Santander",
                "category": expense,
                "tags": "Viaje Mty",
            },
            {
                "block": "expense",
                "amount": 50,
                "description": "Taxi",
                "date": "2026-10-13",
                "account": "Santander",
                "category": expense,
                "tags": "Viaje Mty",
            },
            {
                "block": "expense",
                "amount": 20,
                "description": "Mal",
                "date": "2026-10-13",
                "account": "NoExiste",
                "category": expense,
            },
        ]
    )
    data = await _preview(client, headers, content)

    assert data["valid_rows"] == 3 and data["error_rows"] == 1
    assert (
        data["income"] == {"count": 1, "total": "1000.00"} or float(data["income"]["total"]) == 1000
    )
    assert data["expense"]["count"] == 2 and float(data["expense"]["total"]) == 150
    assert (data["date_from"], data["date_to"]) == ("2026-10-01", "2026-10-13")
    assert [(a["name"], a["count"]) for a in data["by_account"]] == [("Santander", 3)]
    assert [(t["name"], t["count"], float(t["total"])) for t in data["by_tag"]] == [
        ("Viaje Mty", 2, 150.0)
    ]
    assert any(c["type"] == "expense" and c["count"] == 2 for c in data["by_category"])
    assert data["errors"][0]["row"] == 6 and "NoExiste" in data["errors"][0]["reason"]
    assert data["possible_duplicates"]["count"] == 0

    # nothing was saved
    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    assert listed.json()["data"] == []


async def test_preview_warns_about_rows_that_look_already_loaded(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    expense = await _get_category_names(client, headers, "expense")
    rows = [
        {
            "block": "expense",
            "amount": 100,
            "description": "Cena",
            "date": "2026-10-12",
            "account": "Santander",
            "category": expense,
        },
        {
            "block": "expense",
            "amount": 60,
            "description": "Nueva",
            "date": "2026-10-12",
            "account": "Santander",
            "category": expense,
        },
    ]
    first = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", _build_upload(rows[:1]), "application/octet-stream")},
    )
    assert first.json()["data"]["created"] == 1

    data = await _preview(client, headers, _build_upload(rows))
    assert data["valid_rows"] == 2
    assert data["possible_duplicates"]["count"] == 1
    example = data["possible_duplicates"]["examples"][0]
    assert example["description"] == "Cena" and float(example["amount"]) == 100


async def test_preview_of_a_file_without_the_sheet_reports_the_error(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    wb = Workbook()
    wb.active.title = "OtraHoja"
    buf = BytesIO()
    wb.save(buf)
    data = await _preview(client, headers, buf.getvalue())
    assert data["valid_rows"] == 0 and data["error_rows"] == 1
    assert "Transacciones" in data["errors"][0]["reason"]


async def test_template_has_the_notes_columns(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    ws = load_workbook(BytesIO(response.content))["Transacciones"]
    assert ws["M2"].value == "Notas"
    assert ws["N2"].value == "Notas (ingresos)"


async def test_upload_saves_the_note_of_each_row(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    expense_category = await _get_category_names(client, headers, "expense")
    income_category = await _get_category_names(client, headers, "income")

    content = _build_upload(
        [
            {
                "block": "expense",
                "amount": 3000,
                "description": "Placas del Lancer (parte 1)",
                "date": "2026-03-05",
                "account": "Santander",
                "category": expense_category,
                "notes": "  Pagadas con la TDC de Jose Manuel  ",
            },
            {
                "block": "expense",
                "amount": 10,
                "description": "Sin nota",
                "date": "2026-03-05",
                "account": "Santander",
                "category": expense_category,
            },
            {
                "block": "expense",
                "amount": 20,
                "description": "Nota en blanco",
                "date": "2026-03-05",
                "account": "Santander",
                "category": expense_category,
                "notes": "   ",
            },
            {
                "block": "income",
                "amount": 2000,
                "description": "Honorarios",
                "date": "2026-03-06",
                "account": "Santander",
                "category": income_category,
                "notes": "Pago semanal",
            },
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    assert response.json()["data"] == {"created": 4, "errors": []}

    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    notes = {t["description"]: t["notes"] for t in listed.json()["data"]}
    assert notes == {
        "Placas del Lancer (parte 1)": "Pagadas con la TDC de Jose Manuel",
        "Sin nota": None,
        "Nota en blanco": None,
        "Honorarios": "Pago semanal",
    }


async def test_a_row_with_an_income_and_an_expense_keeps_each_note_in_its_block(
    client: AsyncClient,
):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    expense_category = await _get_category_names(client, headers, "expense")
    income_category = await _get_category_names(client, headers, "income")

    wb = Workbook()
    ws = wb.active
    ws.title = "Transacciones"
    ws.append([])
    ws.append(["Monto"])
    ws.append(
        [500, "Cobro", "2026-03-06", "Santander", income_category, None]
        + [80, "Comida", "2026-03-06", "Santander", expense_category, None]
        + ["nota del gasto", "nota del ingreso"]
    )
    buf = BytesIO()
    wb.save(buf)
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", buf.getvalue(), "application/octet-stream")},
    )
    assert response.json()["data"] == {"created": 2, "errors": []}

    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    notes = {t["description"]: t["notes"] for t in listed.json()["data"]}
    assert notes == {"Cobro": "nota del ingreso", "Comida": "nota del gasto"}


async def _create_card(client: AsyncClient, headers: dict, initial_balance: str) -> str:
    return await _create_account(
        client,
        headers,
        name="Stori",
        type="liability",
        subtype="credit_card",
        initial_balance=initial_balance,
    )


async def _get_balance(client: AsyncClient, headers: dict, account_id: str) -> str:
    response = await client.get(f"/api/v1/accounts/{account_id}", headers=headers)
    return response.json()["data"]["balance"]


async def test_template_has_the_card_payments_block(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    await _create_card(client, headers, "0")

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    wb = load_workbook(BytesIO(response.content))
    ws = wb["Transacciones"]
    assert [ws.cell(row=2, column=c).value for c in range(16, 21)] == [
        "Monto",
        "Descripción",
        "Fecha",
        "Cuenta origen",
        "Tarjeta",
    ]
    assert ws["P1"].value == "Pagos de tarjeta"
    # source accounts (only money accounts) and cards, as dropdown lists
    assert wb["Listas"]["E1"].value == "Santander"
    assert wb["Listas"]["F1"].value == "Stori"


async def test_upload_registers_card_payments_as_transfers(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    bank = await _create_account(client, headers, initial_balance="5000")
    card = await _create_account(
        client,
        headers,
        name="Stori",
        type="liability",
        subtype="credit_card",
        initial_balance="1000",
    )

    content = _build_upload(
        [
            {
                "block": "card_payment",
                "amount": 240,
                "description": "Pago a Stori",
                "date": "2026-08-10",
                "account": "Santander",
                "category": "Stori",
            },
            {
                "block": "card_payment",
                "amount": 150,
                "description": "",
                "date": "2026-08-12",
                "account": "Santander",
                "category": "stori",
            },
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    assert response.json()["data"] == {"created": 2, "errors": []}

    # the money leaves the bank and the card owes less; nothing counts as income/expense
    assert Decimal(await _get_balance(client, headers, bank)) == Decimal("4610")
    assert Decimal(await _get_balance(client, headers, card)) == Decimal("610")
    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    rows = {t["date"]: t for t in listed.json()["data"]}
    assert {t["entry_type"] for t in rows.values()} == {"transfer"}
    assert rows["2026-08-12"]["description"] == "Pago de tarjeta"


async def test_card_payment_rows_with_a_wrong_source_or_card_are_rejected(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    await _create_account(
        client,
        headers,
        name="Stori",
        type="liability",
        subtype="credit_card",
        initial_balance="500",
    )

    def payment(**overrides):
        row = {
            "block": "card_payment",
            "amount": 100,
            "description": "Pago",
            "date": "2026-08-10",
            "account": "Santander",
            "category": "Stori",
        }
        return {**row, **overrides}

    content = _build_upload(
        [
            payment(category="Santander"),  # the "card" is a bank account
            payment(account="Stori"),  # the source is a credit card
            payment(category="NoExiste"),
            payment(),  # this one is fine
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    data = response.json()["data"]
    assert data["created"] == 1
    assert [e["row"] for e in data["errors"]] == [3, 4, 5]
    assert all(e["block"] == "Pago de tarjeta" for e in data["errors"])


async def test_preview_counts_card_payments_apart_from_income_and_expenses(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    bank = await _create_account(client, headers, initial_balance="5000")
    card = await _create_account(
        client,
        headers,
        name="Stori",
        type="liability",
        subtype="credit_card",
        initial_balance="1000",
    )
    expense = await _get_category_names(client, headers, "expense")

    content = _build_upload(
        [
            {
                "block": "expense",
                "amount": 80,
                "description": "Comida",
                "date": "2026-08-10",
                "account": "Santander",
                "category": expense,
            },
            {
                "block": "card_payment",
                "amount": 240,
                "description": "Pago a Stori",
                "date": "2026-08-10",
                "account": "Santander",
                "category": "Stori",
            },
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/preview",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    data = response.json()["data"]
    assert data["valid_rows"] == 2
    assert data["expense"] == {"count": 1, "total": "80"}
    assert data["income"] == {"count": 0, "total": "0"}
    assert data["card_payment"] == {"count": 1, "total": "240"}
    assert [c["name"] for c in data["by_category"]] == [expense]
    # nothing was saved by the preview
    assert Decimal(await _get_balance(client, headers, bank)) == Decimal("5000")
    assert Decimal(await _get_balance(client, headers, card)) == Decimal("1000")

    # after loading, the same file warns that the payment looks already loaded
    await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    again = await client.post(
        "/api/v1/bulk-import/preview",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    assert again.json()["data"]["possible_duplicates"]["count"] == 2


async def test_template_has_the_installments_column(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)

    response = await client.get("/api/v1/bulk-import/template", headers=headers)
    ws = load_workbook(BytesIO(response.content))["Transacciones"]
    assert ws["U2"].value == "Meses sin intereses"


async def test_upload_creates_installment_plans_for_card_purchases(client: AsyncClient):
    token = await _register_and_login(client)
    headers = {"Authorization": f"Bearer {token}"}
    await _create_account(client, headers)
    await _create_card(client, headers, "0")
    expense = await _get_category_names(client, headers, "expense")

    def purchase(description: str, account: str, **extra):
        return {
            "block": "expense",
            "amount": 2045.67,
            "description": description,
            "date": "2026-07-10",
            "account": account,
            "category": expense,
            **extra,
        }

    content = _build_upload(
        [
            purchase("Teclado a 12 meses", "Stori", months=12),
            purchase("Compra normal con tarjeta", "Stori"),
            purchase("Meses en una cuenta de dinero", "Santander", months=6),
            purchase("Meses invalidos", "Stori", months="uno"),
            purchase("Meses de menos", "Stori", months=1),
        ]
    )
    response = await client.post(
        "/api/v1/bulk-import/upload",
        headers=headers,
        files={"file": ("lote.xlsx", content, "application/octet-stream")},
    )
    data = response.json()["data"]
    assert data["created"] == 2
    assert sorted(e["row"] for e in data["errors"]) == [5, 6, 7]

    listed = await client.get("/api/v1/transactions?per_page=50", headers=headers)
    plans = {t["description"]: t["installment"] for t in listed.json()["data"]}
    assert plans["Compra normal con tarjeta"] is None
    assert plans["Teclado a 12 meses"]["total_installments"] == 12
