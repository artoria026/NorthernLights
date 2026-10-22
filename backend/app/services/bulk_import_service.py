import re
from dataclasses import dataclass, field
from datetime import date as date_type
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.account import Account
from app.models.category import Category
from app.models.debt import Debt, DebtMovement
from app.models.recurring import RecurringItem
from app.models.tag import Tag
from app.models.transaction import JournalEntry
from app.schemas.debt import DebtCreate, DebtIncreaseCreate, DebtPaymentCreate, DebtUpdate
from app.schemas.recurring import RecurringItemCreate
from app.schemas.transaction import JournalLineIn, TransactionCreate
from app.services import (
    account_service,
    category_service,
    debt_service,
    recurring_service,
    transaction_service,
)

# Layout: two blocks side by side on a single sheet, like a typical
# homemade spreadsheet (see the example the user brought) -- Income in A-E,
# empty column F as a separator, Expenses in G-K. Same 5 columns in both:
# Amount, Description, Date, Account, Category (the last two are the ones
# we added to their original format -- the app needs them to know where
# each transaction goes).
SHEET_NAME = "Transacciones"
LISTS_SHEET = "Listas"
TEMPLATE_ROWS = 200
HEADER_ROW = 2
FIRST_DATA_ROW = 3

BLOCK_HEADERS = ["Monto", "Descripción", "Fecha", "Cuenta", "Categoría"]
INCOME_START_COL = 1  # A
EXPENSE_START_COL = 7  # G

# Optional third block (P-T): payments to a credit card, a transfer from one of the user's
# own accounts to the card. It has no category -- the last column is the card instead.
CARD_PAYMENT_START_COL = 16  # P
CARD_PAYMENT_HEADERS = ["Monto", "Descripción", "Fecha", "Cuenta origen", "Tarjeta"]
CARD_PAYMENT_DEFAULT_DESCRIPTION = "Pago de tarjeta"
CARD_PAYMENT_FILL = PatternFill("solid", fgColor="D8E4BC")

# Optional column after the Expenses block (L): the tags of each expense, separated by
# commas. Files made before it existed simply don't have it -- they keep working.
TAGS_COL = 12  # L
TAGS_HEADER = "Etiquetas"

# Two more optional columns after the tags: a free-text note per transaction, one for
# the Expenses block (M, next to its tags) and one for the Income block (N), because a
# single row can hold an income and an expense at once. Old files don't have them.
EXPENSE_NOTES_COL = 13  # M
INCOME_NOTES_COL = 14  # N
NOTES_HEADER = "Notas"
INCOME_NOTES_HEADER = "Notas (ingresos)"

# Optional column after the card payments block (U): months without interest for an expense
# paid with a credit card ("MSI"). Empty = a normal purchase; 2 or more = an installment plan.
INSTALLMENTS_COL = 21  # U
INSTALLMENTS_HEADER = "Meses sin intereses"
MIN_INSTALLMENTS = 2
TAG_SEPARATORS = re.compile(r"[,;]")

# Optional block after the installments column (W-AC): the loans and repayments of the people
# who owe you (or that you owe). One row per step: a loan opens the person's debt or adds to
# it, a repayment lowers it, an opening balance registers what was already owed before the
# app. The last column (direction) is optional and defaults to "Me deben".
DEBT_START_COL = 23  # W
DEBT_HEADERS = ["Monto", "Persona", "Movimiento", "Fecha", "Cuenta", "Concepto"]
DEBT_DIRECTION_COL = 29  # AC
DEBT_DIRECTION_HEADER = "Dirección"
# Optional payment plan of the debt (AD-AH), read from any row of that person: the periodic
# payment, how often, the day it falls on, the next due date and the account it is paid from.
DEBT_PLAN_AMOUNT_COL = 30  # AD
DEBT_PLAN_FREQUENCY_COL = 31  # AE
DEBT_PLAN_DAY_COL = 32  # AF
DEBT_PLAN_NEXT_COL = 33  # AG
DEBT_PLAN_ACCOUNT_COL = 34  # AH
DEBT_PLAN_HEADERS = {
    DEBT_PLAN_AMOUNT_COL: "Pago periódico",
    DEBT_PLAN_FREQUENCY_COL: "Frecuencia",
    DEBT_PLAN_DAY_COL: "Día de pago",
    DEBT_PLAN_NEXT_COL: "Siguiente pago",
    DEBT_PLAN_ACCOUNT_COL: "Cuenta de pago",
}
DEBT_FREQUENCY_OPTIONS = ["Mensual", "Quincenal", "Semanal"]
# Optional kind of the debt (AI), read when the row creates it. Empty = informal, what a
# person's loan has always been in this block.
DEBT_TYPE_COL = 35  # AI
DEBT_TYPE_HEADER = "Tipo"
DEBT_TYPE_OPTIONS = [
    "Préstamo informal",
    "Préstamo personal",
    "Crédito de nómina",
    "Crédito cívico",
    "Préstamo recibido",
]
_DEBT_TYPES = {
    "": "informal",
    "préstamo informal": "informal",
    "prestamo informal": "informal",
    "informal": "informal",
    "préstamo personal": "personal_loan",
    "prestamo personal": "personal_loan",
    "personal_loan": "personal_loan",
    "crédito de nómina": "payroll_loan",
    "credito de nomina": "payroll_loan",
    "payroll_loan": "payroll_loan",
    "crédito cívico": "civic",
    "credito civico": "civic",
    "civic": "civic",
    "préstamo recibido": "loan_received",
    "prestamo recibido": "loan_received",
    "loan_received": "loan_received",
}
_DEBT_FREQUENCIES = {
    "mensual": "monthly",
    "monthly": "monthly",
    "quincenal": "biweekly",
    "biweekly": "biweekly",
    "semanal": "weekly",
    "weekly": "weekly",
}
DEBT_MOVEMENT_OPTIONS = ["Préstamo", "Abono", "Saldo inicial"]
DEBT_DIRECTION_OPTIONS = ["Me deben", "Debo"]
DEBT_FILL = PatternFill("solid", fgColor="FFE699")
_DEBT_MOVEMENTS = {
    "prestamo": "loan",
    "préstamo": "loan",
    "abono": "payment",
    "pago": "payment",
    "cobro": "payment",
    "saldo inicial": "opening",
    "saldo": "opening",
}
_DEBT_DIRECTIONS = {
    "": "owed_to_me",
    "me deben": "owed_to_me",
    "owed_to_me": "owed_to_me",
    "debo": "owed_by_me",
    "owed_by_me": "owed_by_me",
}

# Optional block after the debts (AK-AS): recurring payments or incomes -- subscriptions, services,
# insurance -- that the app will keep charging by itself. One row per item. The first 5 columns
# follow the usual block order; the rest are optional (type defaults to subscription).
RECURRING_START_COL = 37  # AK
RECURRING_HEADERS = ["Nombre", "Monto", "Frecuencia", "Siguiente cobro", "Cuenta"]
RECURRING_CATEGORY_COL = 42  # AP
RECURRING_CATEGORY_HEADER = "Categoría"
RECURRING_TYPE_COL = 43  # AQ
RECURRING_TYPE_HEADER = "Tipo"
RECURRING_DAY_COL = 44  # AR
RECURRING_DAY_HEADER = "Día del mes"
RECURRING_NOTES_COL = 45  # AS
RECURRING_NOTES_HEADER = "Notas"
RECURRING_FREQUENCY_OPTIONS = ["Mensual", "Semanal", "Quincenal", "Bimestral", "Anual"]
RECURRING_TYPE_OPTIONS = ["Suscripción", "Servicio", "Servicio público", "Ingreso"]
_RECURRING_FREQUENCIES = {
    "semanal": "weekly",
    "weekly": "weekly",
    "quincenal": "biweekly",
    "biweekly": "biweekly",
    "mensual": "monthly",
    "monthly": "monthly",
    "bimestral": "bimonthly",
    "bimonthly": "bimonthly",
    "anual": "annual",
    "annual": "annual",
}
_RECURRING_TYPES = {
    "": "subscription",
    "suscripción": "subscription",
    "suscripcion": "subscription",
    "subscription": "subscription",
    "servicio": "service",
    "service": "service",
    "servicio público": "utility",
    "servicio publico": "utility",
    "utility": "utility",
    "ingreso": "income",
    "income": "income",
}
RECURRING_FILL = PatternFill("solid", fgColor="C6E0B4")
# Two more optional columns after the recurring block: the NAME of the recurring item a transaction
# pays or collects (one for the Expenses block, one for the Income block, like the notes). The
# transaction then carries the link, as the ones the app generates by itself. The item must be
# active, or be created in the same file (Recurrentes block).
EXPENSE_RECURRING_COL = 46  # AT
INCOME_RECURRING_COL = 47  # AU
EXPENSE_RECURRING_HEADER = "Recurrente (gastos)"
INCOME_RECURRING_HEADER = "Recurrente (ingresos)"

INCOME_FILL = PatternFill("solid", fgColor="B8CCE4")
EXPENSE_FILL = PatternFill("solid", fgColor="E6B8B7")

COLUMN_WIDTHS = {
    1: 12,
    2: 30,
    3: 14,
    4: 20,
    5: 24,
    6: 3,
    7: 12,
    8: 30,
    9: 14,
    10: 20,
    11: 24,
    12: 28,
    13: 40,
    14: 40,
    15: 3,
    16: 12,
    17: 30,
    18: 14,
    19: 20,
    20: 24,
    21: 18,
    22: 3,
    23: 12,
    24: 24,
    25: 16,
    26: 14,
    27: 20,
    28: 36,
    29: 14,
    30: 14,
    31: 14,
    32: 12,
    33: 14,
    34: 20,
    35: 20,
    36: 3,
    37: 30,
    38: 12,
    39: 14,
    40: 16,
    41: 20,
    42: 24,
    43: 18,
    44: 12,
    45: 36,
    46: 30,
    47: 30,
}


def _write_block(
    ws, start_col: int, title: str, fill: PatternFill, headers: list[str] = BLOCK_HEADERS
) -> None:
    title_cell = ws.cell(row=1, column=start_col, value=title)
    title_cell.font = Font(bold=True)
    for offset in range(5):
        ws.cell(row=1, column=start_col + offset).fill = fill
        header_cell = ws.cell(row=HEADER_ROW, column=start_col + offset, value=headers[offset])
        header_cell.font = Font(bold=True)
        header_cell.fill = fill


def _is_asset(account: Account) -> bool:
    return account.type == "asset"


def _is_credit_card(account: Account) -> bool:
    return account.type == "liability" and account.subtype == "credit_card"


def build_template_workbook(
    accounts: list[Account],
    income_categories: list[Category],
    expense_categories: list[Category],
    tag_names: list[str] | None = None,
) -> BytesIO:
    """Downloadable template: two tables (Income/Expenses) with Account and
    Category as dropdown lists (referencing a hidden 'Listas' sheet
    -- Excel doesn't support long inline lists beyond ~255 characters). The
    Expenses block has an extra optional 'Etiquetas' column (several tags separated
    by commas; the dropdown suggests the user's tags but doesn't force a single one)."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    _write_block(ws, INCOME_START_COL, "Ingresos", INCOME_FILL)
    _write_block(ws, EXPENSE_START_COL, "Gastos", EXPENSE_FILL)
    _write_block(
        ws, CARD_PAYMENT_START_COL, "Pagos de tarjeta", CARD_PAYMENT_FILL, CARD_PAYMENT_HEADERS
    )
    _write_block(ws, DEBT_START_COL, "Deudas", DEBT_FILL, DEBT_HEADERS[:5])
    # The debts block has a sixth column (concept) and the optional direction after it.
    for col, header in (
        (DEBT_START_COL + 5, DEBT_HEADERS[5]),
        (DEBT_DIRECTION_COL, DEBT_DIRECTION_HEADER),
    ):
        debt_header = ws.cell(row=HEADER_ROW, column=col, value=header)
        debt_header.font = Font(bold=True)
        debt_header.fill = DEBT_FILL
        ws.cell(row=1, column=col).fill = DEBT_FILL
    for col, header in (*DEBT_PLAN_HEADERS.items(), (DEBT_TYPE_COL, DEBT_TYPE_HEADER)):
        plan_header = ws.cell(row=HEADER_ROW, column=col, value=header)
        plan_header.font = Font(bold=True)
        plan_header.fill = DEBT_FILL
        ws.cell(row=1, column=col).fill = DEBT_FILL
    _write_block(ws, RECURRING_START_COL, "Recurrentes", RECURRING_FILL, RECURRING_HEADERS)
    for col, header in (
        (RECURRING_CATEGORY_COL, RECURRING_CATEGORY_HEADER),
        (RECURRING_TYPE_COL, RECURRING_TYPE_HEADER),
        (RECURRING_DAY_COL, RECURRING_DAY_HEADER),
        (RECURRING_NOTES_COL, RECURRING_NOTES_HEADER),
    ):
        rec_header = ws.cell(row=HEADER_ROW, column=col, value=header)
        rec_header.font = Font(bold=True)
        rec_header.fill = RECURRING_FILL
        ws.cell(row=1, column=col).fill = RECURRING_FILL
    for col, header, fill in (
        (EXPENSE_RECURRING_COL, EXPENSE_RECURRING_HEADER, EXPENSE_FILL),
        (INCOME_RECURRING_COL, INCOME_RECURRING_HEADER, INCOME_FILL),
    ):
        link_header = ws.cell(row=HEADER_ROW, column=col, value=header)
        link_header.font = Font(bold=True)
        link_header.fill = fill
        ws.cell(row=1, column=col).fill = fill
    tags_header = ws.cell(row=HEADER_ROW, column=TAGS_COL, value=TAGS_HEADER)
    tags_header.font = Font(bold=True)
    tags_header.fill = EXPENSE_FILL
    ws.cell(row=1, column=TAGS_COL).fill = EXPENSE_FILL
    for col, header, fill in (
        (EXPENSE_NOTES_COL, NOTES_HEADER, EXPENSE_FILL),
        (INCOME_NOTES_COL, INCOME_NOTES_HEADER, INCOME_FILL),
        (INSTALLMENTS_COL, INSTALLMENTS_HEADER, EXPENSE_FILL),
    ):
        notes_header = ws.cell(row=HEADER_ROW, column=col, value=header)
        notes_header.font = Font(bold=True)
        notes_header.fill = fill
        ws.cell(row=1, column=col).fill = fill

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A3"

    lists_ws = wb.create_sheet(LISTS_SHEET)
    lists_ws.sheet_state = "hidden"
    account_names = [a.name for a in accounts] or ["(crea una cuenta primero)"]
    income_names = [c.display_name for c in income_categories] or ["(sin categorias)"]
    expense_names = [c.display_name for c in expense_categories] or ["(sin categorias)"]
    for i, name in enumerate(account_names, start=1):
        lists_ws.cell(row=i, column=1, value=name)
    for i, name in enumerate(income_names, start=1):
        lists_ws.cell(row=i, column=2, value=name)
    for i, name in enumerate(expense_names, start=1):
        lists_ws.cell(row=i, column=3, value=name)
    tag_names = list(tag_names or [])
    for i, name in enumerate(tag_names, start=1):
        lists_ws.cell(row=i, column=4, value=name)
    # For the card payments block: where the money comes from (assets) and the cards.
    source_names = [a.name for a in accounts if _is_asset(a)] or ["(crea una cuenta primero)"]
    card_names = [a.name for a in accounts if _is_credit_card(a)] or ["(sin tarjetas)"]
    for i, name in enumerate(source_names, start=1):
        lists_ws.cell(row=i, column=5, value=name)
    for i, name in enumerate(card_names, start=1):
        lists_ws.cell(row=i, column=6, value=name)

    def _range(col: str, count: int) -> str:
        return f"={LISTS_SHEET}!${col}$1:${col}${count}"

    last_row = FIRST_DATA_ROW + TEMPLATE_ROWS - 1

    dv_account_income = DataValidation(
        type="list", formula1=_range("A", len(account_names)), allow_blank=True
    )
    dv_cat_income = DataValidation(
        type="list", formula1=_range("B", len(income_names)), allow_blank=True
    )
    dv_account_expense = DataValidation(
        type="list", formula1=_range("A", len(account_names)), allow_blank=True
    )
    dv_cat_expense = DataValidation(
        type="list", formula1=_range("C", len(expense_names)), allow_blank=True
    )
    dv_source = DataValidation(
        type="list", formula1=_range("E", len(source_names)), allow_blank=True
    )
    dv_card = DataValidation(type="list", formula1=_range("F", len(card_names)), allow_blank=True)
    dv_debt_account = DataValidation(
        type="list", formula1=_range("A", len(account_names)), allow_blank=True
    )
    dv_debt_movement = DataValidation(
        type="list", formula1='"' + ",".join(DEBT_MOVEMENT_OPTIONS) + '"', allow_blank=True
    )
    dv_debt_direction = DataValidation(
        type="list", formula1='"' + ",".join(DEBT_DIRECTION_OPTIONS) + '"', allow_blank=True
    )
    dv_debt_plan_account = DataValidation(
        type="list", formula1=_range("E", len(source_names)), allow_blank=True
    )
    dv_debt_type = DataValidation(
        type="list", formula1='"' + ",".join(DEBT_TYPE_OPTIONS) + '"', allow_blank=True
    )
    dv_rec_frequency = DataValidation(
        type="list", formula1='"' + ",".join(RECURRING_FREQUENCY_OPTIONS) + '"', allow_blank=True
    )
    dv_rec_type = DataValidation(
        type="list", formula1='"' + ",".join(RECURRING_TYPE_OPTIONS) + '"', allow_blank=True
    )
    dv_rec_account = DataValidation(
        type="list", formula1=_range("A", len(account_names)), allow_blank=True
    )
    dv_rec_category = DataValidation(
        type="list", formula1=_range("C", len(expense_names)), allow_blank=True
    )
    dv_debt_frequency = DataValidation(
        type="list", formula1='"' + ",".join(DEBT_FREQUENCY_OPTIONS) + '"', allow_blank=True
    )
    for dv in (
        dv_rec_frequency,
        dv_rec_type,
        dv_rec_account,
        dv_rec_category,
        dv_debt_type,
        dv_debt_plan_account,
        dv_debt_frequency,
        dv_account_income,
        dv_cat_income,
        dv_account_expense,
        dv_cat_expense,
        dv_source,
        dv_card,
        dv_debt_account,
        dv_debt_movement,
        dv_debt_direction,
    ):
        ws.add_data_validation(dv)
    if tag_names:
        # Suggests the existing tags, but typing several separated by commas must not be rejected.
        dv_tags = DataValidation(
            type="list",
            formula1=_range("D", len(tag_names)),
            allow_blank=True,
            showErrorMessage=False,
        )
        ws.add_data_validation(dv_tags)
        dv_tags.add(f"L{FIRST_DATA_ROW}:L{last_row}")
    dv_account_income.add(f"D{FIRST_DATA_ROW}:D{last_row}")
    dv_cat_income.add(f"E{FIRST_DATA_ROW}:E{last_row}")
    dv_account_expense.add(f"J{FIRST_DATA_ROW}:J{last_row}")
    dv_cat_expense.add(f"K{FIRST_DATA_ROW}:K{last_row}")
    dv_source.add(f"S{FIRST_DATA_ROW}:S{last_row}")
    dv_card.add(f"T{FIRST_DATA_ROW}:T{last_row}")
    dv_debt_movement.add(f"Y{FIRST_DATA_ROW}:Y{last_row}")
    dv_debt_account.add(f"AA{FIRST_DATA_ROW}:AA{last_row}")
    dv_debt_direction.add(f"AC{FIRST_DATA_ROW}:AC{last_row}")
    dv_debt_frequency.add(f"AE{FIRST_DATA_ROW}:AE{last_row}")
    dv_debt_type.add(f"AI{FIRST_DATA_ROW}:AI{last_row}")
    dv_rec_frequency.add(f"AM{FIRST_DATA_ROW}:AM{last_row}")
    dv_rec_account.add(f"AO{FIRST_DATA_ROW}:AO{last_row}")
    dv_rec_category.add(f"AP{FIRST_DATA_ROW}:AP{last_row}")
    dv_rec_type.add(f"AQ{FIRST_DATA_ROW}:AQ{last_row}")
    dv_debt_plan_account.add(f"AH{FIRST_DATA_ROW}:AH{last_row}")

    for row in range(FIRST_DATA_ROW, last_row + 1):
        ws.cell(row=row, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=9).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=CARD_PAYMENT_START_COL + 2).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=DEBT_START_COL + 3).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=DEBT_PLAN_NEXT_COL).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=RECURRING_START_COL + 3).number_format = "yyyy-mm-dd"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@dataclass
class ParsedRow:
    entry_type: str
    amount: Decimal
    description: str
    date: date_type
    account_name: str
    category_name: str
    row_number: int
    tags: list[str] = field(default_factory=list)
    notes: str | None = None
    # Only for entry_type 'card_payment': the card that receives the payment (account_name
    # is where the money comes from, category_name stays empty).
    card_name: str = ""
    # Only for an expense paid with a credit card: number of months without interest.
    installments: int | None = None
    # Only for entry_type 'debt': whose debt (description holds the person), what step it is
    # ('loan' | 'payment' | 'opening') and whose side ('owed_to_me' | 'owed_by_me').
    debt_kind: str = ""
    debt_direction: str = "owed_to_me"
    # Only for entry_type 'debt', all optional: the payment plan of that person's debt.
    plan_amount: Decimal | None = None
    plan_frequency: str | None = None
    plan_day: int | None = None
    plan_next_date: date_type | None = None
    plan_account_name: str = ""
    # Only for entry_type 'debt': the kind of the debt the row creates (informal by default).
    debt_type: str = "informal"
    # Only for entry_type 'recurring' (description = name, date = next charge): how often, what kind
    # (subscription | service | utility | income) and the optional day of the month.
    rec_frequency: str = ""
    rec_type: str = "subscription"
    rec_day: int | None = None
    # Only for income/expense rows: name of the recurring item this transaction pays or collects.
    recurring_link: str | None = None


def _parse_notes(row_cells: tuple, entry_type: str) -> str | None:
    """The note of the row's block, or None when the cell is empty or the file is older
    than the notes columns."""
    col = EXPENSE_NOTES_COL if entry_type == "expense" else INCOME_NOTES_COL
    if len(row_cells) < col:
        return None
    value = row_cells[col - 1].value
    text = str(value).strip() if value is not None else ""
    return text or None


def _parse_recurring_link(row_cells: tuple, entry_type: str) -> str | None:
    """Name of the recurring item the row's block points to, or None (also in older files)."""
    col = EXPENSE_RECURRING_COL if entry_type == "expense" else INCOME_RECURRING_COL
    if len(row_cells) < col:
        return None
    value = row_cells[col - 1].value
    text = str(value).strip() if value is not None else ""
    return text or None


def _parse_installments(row_cells: tuple, entry_type: str) -> tuple[int | None, bool]:
    """(months, is_valid). Only expenses can have months; an empty cell (or an old file
    without the column) means a normal purchase."""
    if entry_type != "expense" or len(row_cells) < INSTALLMENTS_COL:
        return None, True
    value = row_cells[INSTALLMENTS_COL - 1].value
    if value is None or str(value).strip() == "":
        return None, True
    try:
        months = int(str(value).strip())
    except ValueError:
        return None, False
    if months < MIN_INSTALLMENTS:
        return None, False
    return months, True


def _parse_tags(value: object) -> list[str]:
    """'Viaje Mty, Choque' -> ['Viaje Mty', 'Choque'] (no blanks, no repeats, order kept)."""
    if value is None:
        return []
    names = [part.strip() for part in TAG_SEPARATORS.split(str(value))]
    return list(dict.fromkeys(name for name in names if name))


@dataclass
class ParseResult:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)


def _parse_date(value: object) -> date_type | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date_type):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()  # noqa: DTZ007 (business date, no time)
            except ValueError:
                continue
    return None


def _parse_amount(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        cleaned = str(value).replace("$", "").replace(",", "").strip()
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None
    return amount if amount > 0 else None


BLOCK_LABELS = {
    "income": "Ingreso",
    "expense": "Gasto",
    "card_payment": "Pago de tarjeta",
    "debt": "Deuda",
    "recurring": "Recurrente",
}


def _parse_block(row_cells: tuple, row_idx: int, entry_type: str, start_col: int) -> tuple:
    """start_col is 0-indexed within `row_cells` (0 for Income/A, 6 for Expenses/G, 15 for
    Card payments/P). Returns (ParsedRow | None, error dict | None)."""
    block_label = BLOCK_LABELS[entry_type]
    is_card_payment = entry_type == "card_payment"
    amount_cell = row_cells[start_col].value
    if amount_cell in (None, ""):
        return None, None

    amount = _parse_amount(amount_cell)
    description = row_cells[start_col + 1].value
    date_value = _parse_date(row_cells[start_col + 2].value)
    account_name = row_cells[start_col + 3].value
    category_name = row_cells[start_col + 4].value
    account_name = str(account_name).strip() if account_name else ""
    category_name = str(category_name).strip() if category_name else ""

    if amount is None:
        return None, {"row": row_idx, "block": block_label, "reason": "Monto inválido"}
    if date_value is None:
        return None, {"row": row_idx, "block": block_label, "reason": "Fecha inválida o vacía"}
    if not account_name:
        reason = "Falta la cuenta de origen" if is_card_payment else "Falta la cuenta"
        return None, {"row": row_idx, "block": block_label, "reason": reason}
    if not category_name:
        reason = "Falta la tarjeta" if is_card_payment else "Falta la categoría"
        return None, {"row": row_idx, "block": block_label, "reason": reason}

    installments, installments_ok = _parse_installments(row_cells, entry_type)
    if not installments_ok:
        return None, {
            "row": row_idx,
            "block": block_label,
            "reason": (
                f"Los meses sin intereses deben ser un número entero de {MIN_INSTALLMENTS} o más"
            ),
        }

    default_description = CARD_PAYMENT_DEFAULT_DESCRIPTION if is_card_payment else block_label
    return (
        ParsedRow(
            entry_type=entry_type,
            amount=amount,
            description=str(description).strip() if description else default_description,
            date=date_value,
            account_name=account_name,
            # In a card payment the last column is the card, not a category.
            category_name="" if is_card_payment else category_name,
            card_name=category_name if is_card_payment else "",
            row_number=row_idx,
            # Only expenses carry tags (the rest of the app enforces it too); the column
            # may not exist in files made before it was added.
            tags=(
                _parse_tags(row_cells[TAGS_COL - 1].value)
                if entry_type == "expense" and len(row_cells) >= TAGS_COL
                else []
            ),
            notes=None if is_card_payment else _parse_notes(row_cells, entry_type),
            installments=installments,
            recurring_link=(
                None if is_card_payment else _parse_recurring_link(row_cells, entry_type)
            ),
        ),
        None,
    )


def _debt_cell(row_cells: tuple, col: int) -> object:
    """Cell of the debts block by 1-based column; None when the file is older than it."""
    return row_cells[col - 1].value if len(row_cells) >= col else None


def _parse_debt_plan(row_cells: tuple) -> tuple[dict, str | None]:
    """The optional payment plan of a debt row (columns AD-AH): ({ParsedRow fields}, error).
    No plan at all is fine; with one, the periodic payment is required and must be positive."""
    amount_cell = _debt_cell(row_cells, DEBT_PLAN_AMOUNT_COL)
    frequency_cell = _debt_cell(row_cells, DEBT_PLAN_FREQUENCY_COL)
    day_cell = _debt_cell(row_cells, DEBT_PLAN_DAY_COL)
    next_cell = _debt_cell(row_cells, DEBT_PLAN_NEXT_COL)
    account_cell = _debt_cell(row_cells, DEBT_PLAN_ACCOUNT_COL)
    cells = (amount_cell, frequency_cell, day_cell, next_cell, account_cell)
    if all(c in (None, "") for c in cells):
        return {}, None
    amount = _parse_amount(amount_cell)
    if amount is None or amount <= 0:
        return {}, "El plan de pagos necesita un Pago periódico válido (mayor que 0)"
    frequency_text = str(frequency_cell or "").strip().lower()
    frequency = _DEBT_FREQUENCIES.get(frequency_text) if frequency_text else "monthly"
    if frequency is None:
        return {}, "Frecuencia inválida: usa Mensual, Quincenal o Semanal"
    day = None
    if day_cell not in (None, ""):
        try:
            day = int(day_cell)
        except (TypeError, ValueError):
            day = None
        if day is None or not 1 <= day <= 31:
            return {}, "El día de pago debe ser un número del 1 al 31"
    next_date = None
    if next_cell not in (None, ""):
        next_date = _parse_date(next_cell)
        if next_date is None:
            return {}, "Fecha inválida en Siguiente pago"
    return {
        "plan_amount": amount,
        "plan_frequency": frequency,
        "plan_day": day,
        "plan_next_date": next_date,
        "plan_account_name": str(account_cell).strip() if account_cell else "",
    }, None


def _parse_debt_row(row_cells: tuple, row_idx: int) -> tuple:
    """One step of a person's debt (columns W-AC). Returns (ParsedRow | None, error | None)."""
    amount_cell = _debt_cell(row_cells, DEBT_START_COL)
    person_cell = _debt_cell(row_cells, DEBT_START_COL + 1)
    if amount_cell in (None, "") and person_cell in (None, ""):
        return None, None

    def error(reason: str) -> tuple:
        return None, {"row": row_idx, "block": BLOCK_LABELS["debt"], "reason": reason}

    person = str(person_cell).strip() if person_cell else ""
    movement_raw = _debt_cell(row_cells, DEBT_START_COL + 2)
    movement = _DEBT_MOVEMENTS.get(str(movement_raw or "").strip().lower())
    amount = _parse_amount(amount_cell)
    date_value = _parse_date(_debt_cell(row_cells, DEBT_START_COL + 3))
    account = _debt_cell(row_cells, DEBT_START_COL + 4)
    concept = _debt_cell(row_cells, DEBT_START_COL + 5)
    direction = _DEBT_DIRECTIONS.get(
        str(_debt_cell(row_cells, DEBT_DIRECTION_COL) or "").strip().lower()
    )

    if not person:
        return error("Falta la persona")
    if movement is None:
        return error("Movimiento inválido: usa Préstamo, Abono o Saldo inicial")
    if amount is None:
        return error("Monto inválido")
    if date_value is None:
        return error("Fecha inválida o vacía")
    if direction is None:
        return error("Dirección inválida: usa Me deben o Debo")
    debt_type = _DEBT_TYPES.get(str(_debt_cell(row_cells, DEBT_TYPE_COL) or "").strip().lower())
    if debt_type is None:
        return error("Tipo inválido: usa " + ", ".join(DEBT_TYPE_OPTIONS))
    plan, plan_error = _parse_debt_plan(row_cells)
    if plan_error:
        return error(plan_error)
    return (
        ParsedRow(
            entry_type="debt",
            amount=amount,
            description=person,
            date=date_value,
            account_name=str(account).strip() if account else "",
            category_name="",
            row_number=row_idx,
            notes=str(concept).strip() if concept else None,
            debt_kind=movement,
            debt_direction=direction,
            debt_type=debt_type,
            **plan,
        ),
        None,
    )


def _parse_recurring_row(row_cells: tuple, row_idx: int) -> tuple:
    """One recurring payment or income (columns AK-AS). Returns (ParsedRow | None, error | None)."""
    name_cell = _debt_cell(row_cells, RECURRING_START_COL)
    amount_cell = _debt_cell(row_cells, RECURRING_START_COL + 1)
    if name_cell in (None, "") and amount_cell in (None, ""):
        return None, None

    def error(reason: str) -> tuple:
        return None, {"row": row_idx, "block": BLOCK_LABELS["recurring"], "reason": reason}

    name = str(name_cell).strip() if name_cell else ""
    amount = _parse_amount(amount_cell)
    frequency = _RECURRING_FREQUENCIES.get(
        str(_debt_cell(row_cells, RECURRING_START_COL + 2) or "").strip().lower()
    )
    next_date = _parse_date(_debt_cell(row_cells, RECURRING_START_COL + 3))
    account = _debt_cell(row_cells, RECURRING_START_COL + 4)
    category = _debt_cell(row_cells, RECURRING_CATEGORY_COL)
    item_type = _RECURRING_TYPES.get(
        str(_debt_cell(row_cells, RECURRING_TYPE_COL) or "").strip().lower()
    )
    day_cell = _debt_cell(row_cells, RECURRING_DAY_COL)
    notes = _debt_cell(row_cells, RECURRING_NOTES_COL)

    if not name:
        return error("Falta el nombre")
    if amount is None or amount <= 0:
        return error("Monto inválido: debe ser mayor que 0")
    if frequency is None:
        return error("Frecuencia inválida: usa Mensual, Semanal, Quincenal, Bimestral o Anual")
    if next_date is None:
        return error("Falta o es inválida la fecha del siguiente cobro")
    if not account:
        return error("Falta la cuenta")
    if not category:
        return error("Falta la categoría")
    if item_type is None:
        return error("Tipo inválido: usa Suscripción, Servicio, Servicio público o Ingreso")
    day = None
    if day_cell not in (None, ""):
        try:
            day = int(day_cell)
        except (TypeError, ValueError):
            day = None
        if day is None or not 0 <= day <= 31:
            return error("El día del mes debe ser un número del 0 al 31")
    return (
        ParsedRow(
            entry_type="recurring",
            amount=amount,
            description=name,
            date=next_date,
            account_name=str(account).strip(),
            category_name=str(category).strip(),
            row_number=row_idx,
            notes=str(notes).strip() if notes else None,
            rec_frequency=frequency,
            rec_type=item_type,
            rec_day=day,
        ),
        None,
    )


def parse_upload(content: bytes) -> ParseResult:
    wb = load_workbook(BytesIO(content), data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        return ParseResult(
            errors=[
                {
                    "row": 0,
                    "block": None,
                    "reason": (
                        f"No encontré la hoja '{SHEET_NAME}'. Usa la plantilla descargada sin "
                        "renombrar la hoja."
                    ),
                }
            ]
        )
    ws = wb[SHEET_NAME]
    result = ParseResult()

    for row_idx, row_cells in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW), start=FIRST_DATA_ROW):
        for entry_type, start_col in (
            ("income", INCOME_START_COL - 1),
            ("expense", EXPENSE_START_COL - 1),
            ("card_payment", CARD_PAYMENT_START_COL - 1),
        ):
            if len(row_cells) < start_col + 5:
                continue
            parsed, error = _parse_block(row_cells, row_idx, entry_type, start_col)
            if parsed:
                result.rows.append(parsed)
            if error:
                result.errors.append(error)
        debt_row, debt_error = _parse_debt_row(row_cells, row_idx)
        if debt_row:
            result.rows.append(debt_row)
        if debt_error:
            result.errors.append(debt_error)
        recurring_row, recurring_error = _parse_recurring_row(row_cells, row_idx)
        if recurring_row:
            result.rows.append(recurring_row)
        if recurring_error:
            result.errors.append(recurring_error)

    return result


def _category_names(category: Category) -> list[str]:
    """Lowercased names a category answers to. A system category answers to its name
    in either language, whatever the language of the template the file came from."""
    return [n.strip().lower() for n in (category.name, category.name_en) if n]


def _category_lookup(categories: list[Category]) -> dict[str, UUID]:
    """Lowercased name -> id (a system category answers to either language)."""
    return {name: c.id for c in categories for name in _category_names(c)}


@dataclass
class ResolvedRow:
    """A parsed row whose account, category and tags exist for this user."""

    row: ParsedRow
    account_id: UUID | None
    account_name: str
    category_id: UUID | None
    category_name: str
    tag_ids: list[UUID]
    tag_names: list[str]
    # Only for a card payment: the card that receives it (account_* is the source).
    card_account_id: UUID | None = None
    card_account_name: str = ""
    # Only for a debt row: what the load will do with it ('create' | 'opening' | 'increase' |
    # 'payment').
    debt_action: str = ""
    # Only for a debt row with a payment plan: the account each periodic payment comes from.
    plan_account_id: UUID | None = None


async def resolve_rows(
    session: AsyncSession, user_id: UUID, parsed: ParseResult
) -> tuple[list[ResolvedRow], list[dict]]:
    """Matches every parsed row against the user's accounts, categories and tags
    WITHOUT writing anything. Rows that don't match go to the error list (with the
    reason); the others come back ready to be created. Both the preview and the real
    load go through here, so what the preview shows is what the load will do."""
    accounts = await account_service.list_accounts(session, user_id)
    income_categories = await category_service.list_categories(session, user_id, "income")
    expense_categories = await category_service.list_categories(session, user_id, "expense")

    account_by_name = {a.name.strip().lower(): a for a in accounts}
    tags = (await session.execute(select(Tag).where(Tag.user_id == user_id))).scalars().all()
    tag_by_name = {t.name.strip().lower(): t for t in tags}
    category_by_type = {
        "income": {n: c for c in income_categories for n in _category_names(c)},
        "expense": {n: c for c in expense_categories for n in _category_names(c)},
    }

    resolved: list[ResolvedRow] = []
    errors = list(parsed.errors)
    seen_recurring: set[tuple[str, UUID]] = set()
    file_recurring = {
        r.description.strip().lower(): r.rec_type
        for r in parsed.rows
        if r.entry_type == "recurring"
    }
    for row in parsed.rows:
        if row.entry_type == "debt":
            continue  # debt steps depend on each other: resolved together, below
        block_label = BLOCK_LABELS[row.entry_type]
        if row.entry_type == "recurring":
            item = await _resolve_recurring_row(
                session, user_id, row, account_by_name, category_by_type, seen_recurring, errors
            )
            if item:
                resolved.append(item)
            continue

        def fail(reason: str, row=row, block_label=block_label) -> None:
            errors.append({"row": row.row_number, "block": block_label, "reason": reason})

        account = account_by_name.get(row.account_name.strip().lower())
        if account is None:
            fail(f"La cuenta '{row.account_name}' no existe")
            continue
        if row.entry_type == "card_payment":
            card = account_by_name.get(row.card_name.strip().lower())
            if card is None or not _is_credit_card(card):
                fail(f"La tarjeta '{row.card_name}' no existe o no es una tarjeta de crédito")
                continue
            if not _is_asset(account):
                fail(f"La cuenta de origen '{row.account_name}' debe ser una cuenta de dinero")
                continue
            resolved.append(
                ResolvedRow(
                    row=row,
                    account_id=account.id,
                    account_name=account.name,
                    category_id=None,
                    category_name="",
                    tag_ids=[],
                    tag_names=[],
                    card_account_id=card.id,
                    card_account_name=card.name,
                )
            )
            continue
        if row.installments and not _is_credit_card(account):
            fail("Los meses sin intereses solo aplican si pagas con una tarjeta de crédito")
            continue
        category = category_by_type[row.entry_type].get(row.category_name.strip().lower())
        if category is None:
            fail(f"La categoría '{row.category_name}' no existe")
            continue
        unknown_tags = [name for name in row.tags if name.lower() not in tag_by_name]
        if unknown_tags:
            fail(
                f"La etiqueta '{unknown_tags[0]}' no existe: créala en Etiquetas "
                "o corrige el nombre"
            )
            continue
        if row.recurring_link:
            link_error = await _check_recurring_link(session, user_id, row, file_recurring)
            if link_error:
                fail(link_error)
                continue
        row_tags = [tag_by_name[name.lower()] for name in row.tags]
        resolved.append(
            ResolvedRow(
                row=row,
                account_id=account.id,
                account_name=account.name,
                category_id=category.id,
                category_name=category.display_name,
                tag_ids=[t.id for t in row_tags],
                tag_names=[t.name for t in row_tags],
            )
        )
    debt_items, debt_errors = await _resolve_debt_rows(
        session,
        user_id,
        [r for r in parsed.rows if r.entry_type == "debt"],
        account_by_name,
    )
    resolved.extend(debt_items)
    errors.extend(debt_errors)
    return resolved, errors


async def _active_recurring(session: AsyncSession, user_id: UUID, name: str) -> list[RecurringItem]:
    """The user's active recurring items called `name` (case-insensitive)."""
    result = await session.execute(
        select(RecurringItem).where(
            RecurringItem.user_id == user_id,
            RecurringItem.status == "active",
            RecurringItem.deleted_at.is_(None),
            func.lower(RecurringItem.name) == name.strip().lower(),
        )
    )
    return list(result.scalars().all())


def _link_matches(item_type: str, entry_type: str) -> bool:
    """A recurring income is paid by incomes only; the rest (subscription, service, utility) by
    expenses only."""
    return (item_type == "income") == (entry_type == "income")


async def _check_recurring_link(
    session: AsyncSession, user_id: UUID, row: ParsedRow, file_recurring: dict[str, str]
) -> str | None:
    """Why the recurring item a transaction points to is not usable, or None when it is."""
    key = row.recurring_link.strip().lower()
    items = await _active_recurring(session, user_id, key)
    if any(_link_matches(i.item_type, row.entry_type) for i in items):
        return None
    if key in file_recurring and _link_matches(file_recurring[key], row.entry_type):
        return None
    if items or key in file_recurring:
        kind = "de ingresos" if row.entry_type != "income" else "de gastos"
        return f"El recurrente '{row.recurring_link}' es {kind}: no corresponde a este movimiento"
    return (
        f"No existe el recurrente '{row.recurring_link}': créalo en el bloque Recurrentes "
        "o en la pantalla de Recurrentes"
    )


async def _find_recurring_id(
    session: AsyncSession, user_id: UUID, row: ParsedRow, account_id: UUID | None
) -> UUID | None:
    """The active recurring item a transaction points to, once the ones of the same file exist:
    the one on the same account when there are several."""
    items = [
        i
        for i in await _active_recurring(session, user_id, row.recurring_link or "")
        if _link_matches(i.item_type, row.entry_type)
    ]
    if not items:
        return None
    same_account = [i for i in items if i.account_id == account_id]
    return (same_account or items)[0].id


async def _resolve_recurring_row(
    session: AsyncSession,
    user_id: UUID,
    row: ParsedRow,
    account_by_name: dict[str, Account],
    category_by_type: dict[str, dict],
    seen: set[tuple[str, UUID]],
    errors: list[dict],
) -> ResolvedRow | None:
    """Checks a recurring item WITHOUT writing: its account and category exist (the category must
    be of the type of entry it generates) and the same item isn't already active."""
    label = BLOCK_LABELS["recurring"]

    def fail(reason: str) -> None:
        errors.append({"row": row.row_number, "block": label, "reason": reason})

    account = account_by_name.get(row.account_name.strip().lower())
    if account is None:
        fail(f"La cuenta '{row.account_name}' no existe")
        return None
    entry_type = "income" if row.rec_type == "income" else "expense"
    category = category_by_type[entry_type].get(row.category_name.strip().lower())
    if category is None:
        fail(
            f"La categoría '{row.category_name}' no existe o no es de "
            f"{'ingresos' if entry_type == 'income' else 'gastos'}"
        )
        return None
    key = (row.description.strip().lower(), account.id)
    if key in seen:
        fail(f"'{row.description}' aparece dos veces para la cuenta {account.name}")
        return None
    already = await session.execute(
        select(RecurringItem.id).where(
            RecurringItem.user_id == user_id,
            RecurringItem.account_id == account.id,
            RecurringItem.status == "active",
            RecurringItem.deleted_at.is_(None),
            func.lower(RecurringItem.name) == row.description.strip().lower(),
        )
    )
    if already.first() is not None:
        fail(f"Ya tienes un recurrente activo '{row.description}' en {account.name}")
        return None
    seen.add(key)
    return ResolvedRow(
        row=row,
        account_id=account.id,
        account_name=account.name,
        category_id=category.id,
        category_name=category.display_name,
        tag_ids=[],
        tag_names=[],
    )


async def _open_debt_states(session: AsyncSession, user_id: UUID) -> dict[tuple[str, str], dict]:
    """(person, direction) -> {'active': bool, 'balance': Decimal} for the debts the user already
    has (one per person: an open one wins over a paid-off one)."""
    debts = (
        (
            await session.execute(
                select(Debt)
                .where(
                    Debt.user_id == user_id,
                    Debt.deleted_at.is_(None),
                    Debt.status.in_(("active", "completed")),
                )
                .order_by(Debt.created_at)
            )
        )
        .scalars()
        .all()
    )
    states: dict[tuple[str, str], dict] = {}
    for debt in debts:
        key = (debt.name.strip().lower(), debt.direction)
        previous = states.get(key)
        if previous and previous["active"] and debt.status != "active":
            continue
        states[key] = {"active": debt.status == "active", "balance": debt.current_balance}
    return states


async def _resolve_debt_rows(
    session: AsyncSession,
    user_id: UUID,
    rows: list[ParsedRow],
    account_by_name: dict[str, Account],
) -> tuple[list[ResolvedRow], list[dict]]:
    """Checks the debt steps WITHOUT writing: in date order (then file order), following each
    person's balance as the file would move it, so a repayment right after its loan in the
    same file is valid and one that goes past what is owed is not."""
    resolved: list[ResolvedRow] = []
    errors: list[dict] = []
    if not rows:
        return resolved, errors
    states = await _open_debt_states(session, user_id)
    label = BLOCK_LABELS["debt"]

    for row in sorted(rows, key=lambda r: (r.date, r.row_number)):

        def fail(reason: str, row=row) -> None:
            errors.append({"row": row.row_number, "block": label, "reason": reason})

        person = row.description
        key = (person.strip().lower(), row.debt_direction)
        state = states.get(key)
        needs_account = row.debt_kind in ("loan", "payment")
        account = (
            account_by_name.get(row.account_name.strip().lower()) if row.account_name else None
        )
        if needs_account and not row.account_name:
            fail("Falta la cuenta: de dónde sale o a dónde entra el dinero")
            continue
        if row.account_name and account is None:
            fail(f"La cuenta '{row.account_name}' no existe")
            continue

        plan_account = None
        if row.plan_account_name:
            plan_account = account_by_name.get(row.plan_account_name.strip().lower())
            if plan_account is None or not _is_asset(plan_account):
                fail(f"La cuenta de pago '{row.plan_account_name}' no existe o no es de dinero")
                continue

        if row.debt_kind == "loan":
            action = "increase" if state else "create"
            states[key] = {
                "active": True,
                "balance": (state["balance"] if state else Decimal(0)) + row.amount,
            }
        elif row.debt_kind == "opening":
            if state and state["active"]:
                fail(f"Ya existe una deuda abierta con {person}: usa Préstamo para sumarle")
                continue
            action = "opening"
            states[key] = {"active": True, "balance": row.amount}
        else:
            if not state or not state["active"] or state["balance"] <= 0:
                fail(f"No hay una deuda abierta con {person} (registra antes su Préstamo)")
                continue
            if row.amount > state["balance"]:
                fail(
                    f"El abono ({row.amount}) es mayor que lo que {person} debe "
                    f"en ese momento ({state['balance']})"
                )
                continue
            action = "payment"
            balance = state["balance"] - row.amount
            states[key] = {"active": balance > 0, "balance": balance}

        resolved.append(
            ResolvedRow(
                row=row,
                account_id=account.id if account else None,
                account_name=account.name if account else "",
                category_id=None,
                category_name="",
                tag_ids=[],
                tag_names=[],
                debt_action=action,
                plan_account_id=plan_account.id if plan_account else None,
            )
        )
    return resolved, errors


async def _find_debt(session: AsyncSession, user_id: UUID, person: str, direction: str) -> Debt:
    """The person's debt as it is right now: the open one, or the last one paid off."""
    debts = (
        (
            await session.execute(
                select(Debt)
                .where(
                    Debt.user_id == user_id,
                    Debt.direction == direction,
                    Debt.deleted_at.is_(None),
                    Debt.status.in_(("active", "completed")),
                    func.lower(Debt.name) == person.strip().lower(),
                )
                .order_by(Debt.created_at.desc())
            )
        )
        .scalars()
        .all()
    )
    open_ones = [d for d in debts if d.status == "active"]
    return (open_ones or debts)[0]


async def _apply_debt_row(session: AsyncSession, user_id: UUID, item: ResolvedRow) -> None:
    row = item.row
    if item.debt_action in ("create", "opening"):
        await debt_service.create_debt(
            session,
            user_id,
            DebtCreate(
                name=row.description,
                type=row.debt_type,
                direction=row.debt_direction,
                total_amount=row.amount,
                funding_account_id=item.account_id if item.debt_action == "create" else None,
                start_date=row.date,
                concept=row.notes,
            ),
            "user",
        )
        return
    debt = await _find_debt(session, user_id, row.description, row.debt_direction)
    if item.debt_action == "increase":
        await debt_service.increase_debt(
            session,
            user_id,
            debt.id,
            DebtIncreaseCreate(
                account_id=item.account_id, amount=row.amount, date=row.date, concept=row.notes
            ),
        )
    else:
        await debt_service.register_debt_payment(
            session,
            user_id,
            debt.id,
            DebtPaymentCreate(
                account_id=item.account_id, amount=row.amount, date=row.date, notes=row.notes
            ),
        )


async def _apply_debt_plan(session: AsyncSession, user_id: UUID, item: ResolvedRow) -> None:
    """Sets the payment plan the row carries on its person's debt."""
    row = item.row
    debt = await _find_debt(session, user_id, row.description, row.debt_direction)
    fields: dict = {"payment_amount": row.plan_amount, "payment_frequency": row.plan_frequency}
    if row.plan_day is not None:
        fields["payment_day"] = row.plan_day
    if row.plan_next_date is not None:
        fields["next_payment_date"] = row.plan_next_date
    if item.plan_account_id is not None:
        fields["payment_source_account_id"] = item.plan_account_id
    await debt_service.update_debt(session, user_id, debt.id, DebtUpdate(**fields))


async def commit_parsed_rows(session: AsyncSession, user_id: UUID, parsed: ParseResult) -> dict:
    resolved, errors = await resolve_rows(session, user_id, parsed)
    # The recurring items go first: the transactions of the same file may point to them.
    resolved.sort(key=lambda item: item.row.entry_type != "recurring")
    created = 0
    # person's payment plan -> the last row of the file that carries one (and was applied).
    plans: dict[tuple[str, str], ResolvedRow] = {}
    for item in resolved:
        row = item.row
        if row.entry_type == "debt":
            try:
                async with session.begin_nested():
                    await _apply_debt_row(session, user_id, item)
                if row.plan_amount is not None:
                    plans[(row.description.strip().lower(), row.debt_direction)] = item
            except HTTPException as exc:
                errors.append(
                    {
                        "row": row.row_number,
                        "block": BLOCK_LABELS["debt"],
                        "reason": str(exc.detail),
                    }
                )
                continue
            created += 1
            continue
        if row.entry_type == "recurring":
            try:
                async with session.begin_nested():
                    await recurring_service.create_recurring_item(
                        session,
                        user_id,
                        RecurringItemCreate(
                            name=row.description,
                            item_type=row.rec_type,
                            amount=row.amount,
                            frequency=row.rec_frequency,
                            frequency_day=row.rec_day,
                            account_id=item.account_id,
                            category_id=item.category_id,
                            next_date=row.date,
                            notes=row.notes,
                        ),
                    )
            except HTTPException as exc:
                errors.append(
                    {
                        "row": row.row_number,
                        "block": BLOCK_LABELS["recurring"],
                        "reason": str(exc.detail),
                    }
                )
                continue
            created += 1
            continue
        created += 1
        if row.entry_type == "card_payment":
            # Paying a card is a transfer: it lowers what the card owes (debit on the
            # card) and takes the money out of the source account (credit).
            await transaction_service.create_transaction(
                session,
                user_id,
                TransactionCreate(
                    date=row.date,
                    description=row.description,
                    entry_type="transfer",
                    lines=[
                        JournalLineIn(
                            account_id=item.card_account_id, amount=row.amount, type="debit"
                        ),
                        JournalLineIn(account_id=item.account_id, amount=row.amount, type="credit"),
                    ],
                ),
            )
            continue
        recurring_id = None
        if row.recurring_link:
            recurring_id = await _find_recurring_id(session, user_id, row, item.account_id)
            if recurring_id is None:
                errors.append(
                    {
                        "row": row.row_number,
                        "block": BLOCK_LABELS[row.entry_type],
                        "reason": f"No encontré el recurrente '{row.recurring_link}'",
                    }
                )
                created -= 1
                continue
        await transaction_service.create_transaction(
            session,
            user_id,
            TransactionCreate(
                date=row.date,
                description=row.description,
                entry_type=row.entry_type,
                category_id=item.category_id,
                account_id=item.account_id,
                amount=row.amount,
                tag_ids=item.tag_ids,
                notes=row.notes,
                installment_total=row.installments,
            ),
            is_recurring=recurring_id is not None,
            recurring_id=recurring_id,
        )
    # The plan goes on after every step: each repayment moves the next due date forward, so
    # setting it first would leave it wherever the last repayment of the file put it.
    for item in plans.values():
        try:
            async with session.begin_nested():
                await _apply_debt_plan(session, user_id, item)
        except HTTPException as exc:
            errors.append(
                {
                    "row": item.row.row_number,
                    "block": BLOCK_LABELS["debt"],
                    "reason": str(exc.detail),
                }
            )
    return {"created": created, "errors": errors}


async def _existing_keys(
    session: AsyncSession, user_id: UUID, resolved: list[ResolvedRow]
) -> set[tuple[date_type, Decimal, str, str]]:
    """(date, amount, description, entry_type) of what the user already has in the
    date range of the file, to warn about rows that look already loaded."""
    if not resolved:
        return set()
    start = min(item.row.date for item in resolved)
    end = max(item.row.date for item in resolved)
    entries = (
        (
            await session.execute(
                select(JournalEntry).where(
                    JournalEntry.user_id == user_id,
                    JournalEntry.deleted_at.is_(None),
                    JournalEntry.date >= start,
                    JournalEntry.date <= end,
                    JournalEntry.entry_type.in_(("income", "expense", "transfer")),
                )
            )
        )
        .scalars()
        .all()
    )
    return {
        (e.date, Decimal(e.amount or 0), (e.description or "").strip().lower(), e.entry_type)
        for e in entries
    }


async def _existing_debt_keys(
    session: AsyncSession, user_id: UUID, resolved: list[ResolvedRow]
) -> set[tuple[str, date_type, Decimal, str]]:
    """(person, date, amount, kind) of the debt steps the user already has in the date range of
    the file, to warn about debt rows that look already loaded."""
    dates = [item.row.date for item in resolved if item.row.entry_type == "debt"]
    if not dates:
        return set()
    rows = (
        await session.execute(
            select(Debt.name, DebtMovement.date, DebtMovement.delta, DebtMovement.kind)
            .join(Debt, Debt.id == DebtMovement.debt_id)
            .where(
                DebtMovement.user_id == user_id,
                DebtMovement.date >= min(dates),
                DebtMovement.date <= max(dates),
            )
        )
    ).all()
    return {(name.strip().lower(), date_, abs(delta), kind) for name, date_, delta, kind in rows}


async def build_preview(session: AsyncSession, user_id: UUID, parsed: ParseResult) -> dict:
    """What loading this file would do, without doing it: totals, breakdowns, the rows
    that would be rejected (with the reason) and the ones that look already loaded."""
    resolved, errors = await resolve_rows(session, user_id, parsed)
    existing = await _existing_keys(session, user_id, resolved)
    existing_debt = await _existing_debt_keys(session, user_id, resolved)
    debt_totals = {"loan": [0, Decimal(0)], "payment": [0, Decimal(0)], "opening": [0, Decimal(0)]}
    debt_people: set[tuple[str, str]] = set()
    recurring_total = [0, Decimal(0)]

    by_type = {
        "income": [0, Decimal(0)],
        "expense": [0, Decimal(0)],
        "card_payment": [0, Decimal(0)],
    }
    by_account: dict[str, dict] = {}
    by_category: dict[tuple[str, str], list] = {}
    by_tag: dict[str, list] = {}
    duplicates: list[dict] = []
    for item in resolved:
        row = item.row
        if row.entry_type == "recurring":
            recurring_total[0] += 1
            recurring_total[1] += row.amount
            continue
        if row.entry_type == "debt":
            debt_totals[row.debt_kind][0] += 1
            debt_totals[row.debt_kind][1] += row.amount
            debt_people.add((row.description.strip().lower(), row.debt_direction))
            if (row.description.strip().lower(), row.date, row.amount, row.debt_kind) in (
                existing_debt
            ):
                duplicates.append(
                    {
                        "row": row.row_number,
                        "date": row.date,
                        "amount": row.amount,
                        "description": f"{row.description} ({row.debt_kind})",
                    }
                )
            continue
        by_type[row.entry_type][0] += 1
        by_type[row.entry_type][1] += row.amount
        # A card payment is a transfer between the user's accounts: it isn't income,
        # expense or part of a category, so it only shows in its own total.
        entry_key = "transfer" if row.entry_type == "card_payment" else row.entry_type
        key = (row.date, row.amount, row.description.strip().lower(), entry_key)
        if key in existing:
            duplicates.append(
                {
                    "row": row.row_number,
                    "date": row.date,
                    "amount": row.amount,
                    "description": row.description,
                }
            )
        if row.entry_type == "card_payment":
            continue
        acct = by_account.setdefault(
            item.account_name,
            {"name": item.account_name, "count": 0, "income": Decimal(0), "expense": Decimal(0)},
        )
        acct["count"] += 1
        acct[row.entry_type] += row.amount
        cat = by_category.setdefault((row.entry_type, item.category_name), [0, Decimal(0)])
        cat[0] += 1
        cat[1] += row.amount
        for name in item.tag_names:
            tag = by_tag.setdefault(name, [0, Decimal(0)])
            tag[0] += 1
            tag[1] += row.amount

    # A recurring item's date is its next charge, not a movement: it stays out of the date range.
    dates = [item.row.date for item in resolved if item.row.entry_type != "recurring"]
    return {
        "valid_rows": len(resolved),
        "error_rows": len(errors),
        "income": {"count": by_type["income"][0], "total": by_type["income"][1]},
        "expense": {"count": by_type["expense"][0], "total": by_type["expense"][1]},
        "card_payment": {"count": by_type["card_payment"][0], "total": by_type["card_payment"][1]},
        "debts": {
            "loans": {"count": debt_totals["loan"][0], "total": debt_totals["loan"][1]},
            "payments": {"count": debt_totals["payment"][0], "total": debt_totals["payment"][1]},
            "openings": {"count": debt_totals["opening"][0], "total": debt_totals["opening"][1]},
            "people": len(debt_people),
        },
        "recurring": {"count": recurring_total[0], "total": recurring_total[1]},
        "date_from": min(dates) if dates else None,
        "date_to": max(dates) if dates else None,
        "by_account": sorted(by_account.values(), key=lambda a: -a["count"]),
        "by_category": [
            {"name": name, "type": type_, "count": v[0], "total": v[1]}
            for (type_, name), v in sorted(by_category.items(), key=lambda kv: -kv[1][1])
        ],
        "by_tag": [
            {"name": name, "count": v[0], "total": v[1]}
            for name, v in sorted(by_tag.items(), key=lambda kv: -kv[1][1])
        ],
        "possible_duplicates": {"count": len(duplicates), "examples": duplicates[:5]},
        "errors": errors,
    }
